"""Apply the consolidated review edits to Jeff Robison's deck.

Input : 00_Original_Deck/Applesoft_Loaded_Dice_Deck_1.pptx (untouched)
Output: 01_Edited_Deck/Applesoft_Loaded_Dice_Deck_1_EDITED.pptx

Edits keep Jeff's layout, fonts and colours: existing text boxes are rewritten in place
(first run's formatting is cloned), and new boxes reuse the deck's palette.
Every edit is listed, with reason and source, in 01_Edited_Deck/CHANGES.md.
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "00_Original_Deck" / "Applesoft_Loaded_Dice_Deck_1.pptx"
DST = ROOT / "01_Edited_Deck" / "Applesoft_Loaded_Dice_Deck_1_EDITED.pptx"

INK, MUTED, GREEN, AMBER, PANEL = "E9E7E3", "9DA39B", "5FCB7E", "E0A73E", "202225"


def shape(slide, name_or_id):
    for sh in slide.shapes:
        if sh.shape_id == name_or_id or sh.name == name_or_id:
            return sh
        if sh.shape_type == 6:
            for sub in sh.shapes:
                if sub.shape_id == name_or_id:
                    return sub
    raise KeyError(name_or_id)


def rewrite(sh, lines):
    """Replace a text frame's paragraphs, cloning the first paragraph/run formatting."""
    txBody = sh.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    p0 = paras[0]
    for p in paras[1:]:
        txBody.remove(p)
    for child in list(p0):
        if child.tag in (qn("a:br"), qn("a:fld")):
            p0.remove(child)
    runs = p0.findall(qn("a:r"))
    if not runs:  # empty box: add a run
        r = p0.makeelement(qn("a:r"), {})
        r.append(r.makeelement(qn("a:t"), {}))
        end = p0.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p0.append(r)
        runs = [r]
    for r in runs[1:]:
        p0.remove(r)
    template = copy.deepcopy(p0)
    runs[0].find(qn("a:t")).text = lines[0]
    prev = p0
    for line in lines[1:]:
        np_ = copy.deepcopy(template)
        np_.find(qn("a:r")).find(qn("a:t")).text = line
        prev.addnext(np_)
        prev = np_


def add_text(slide, x, y, w, h, lines, size, color, bold=False, font="Arial", mono=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = "Courier New" if mono else font
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def add_panel(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(PANEL)
    s.line.fill.background()
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


prs = Presentation(SRC)
S = list(prs.slides)
prs.core_properties.title = "Applesoft's Loaded Dice (edited)"
prs.core_properties.subject = "VCF Midwest 21, 2026 - Jeff Robison"

# ---------------------------------------------------------------- slide 1 (badge 0)
notes(S[0], """
Name and title. Fifteen seconds. No bio slide.

BEFORE THIS SLIDE: power-cycle the machine (power switch off, wait, power on), then PRINT RND(1). Do it three times. Do NOT use Ctrl-RESET or PR#6: they keep the fifth seed byte from the last run, so the number changes (Aldridge 1987; Sander-Cederlof, AAL May 1984).

Rehearse on the exact machine you are bringing. Your machine may not print .973136996. What matters is that it prints the same number every power-on.

Applesoft is Microsoft's 6502 BASIC, licensed by Apple. Keep that in mind; it matters on slide 5.
""")

# ---------------------------------------------------------------- slide 2 (badge 1)
add_text(S[1], 0.7, 4.65, 11.9, 0.5,
         ["Rerun the program: different numbers. Power-cycle: the same ones again."],
         18, GREEN, bold=True)
add_text(S[1], 0.7, 5.35, 11.9, 0.8,
         ["“exactly the same sequence each time the machine is powered on.”",
          "J. W. Aldridge, Behavior Research Methods, Instruments, & Computers 19(4), 1987"],
         15, MUTED)
notes(S[1], """
This is the number they just watched three times. Don't explain it yet. "Hold onto that number. By slide 5 you'll know exactly where it lives."

Why it hid for years: rerunning a program, or even rebooting DOS, gives different numbers. Only a power cycle repeats. A programmer at a desk never sees it; the first user of the day does.

If someone says their machine gives a different number: one of the five seed bytes is never initialized (Sander-Cederlof, AAL May 1984), so the exact value can vary by machine. It is still the same on every power-on of that machine.
""")

# ---------------------------------------------------------------- slide 3 (badge 2)
rewrite(shape(S[2], "Text 2"), ["Pseudo-random means repeatable"])
rewrite(shape(S[2], 12), ["Same starting number → same “random” numbers. Entropy only picks the start."])
notes(S[2], """
Thirty seconds. This room knows what a PRNG is.

The only line that matters is the amber one: same starting number, same numbers. Integer BASIC (next slide) depends on it.

If asked: a random seed makes the starting point unpredictable. It does not make a weak generator's output pattern-free. This talk is mostly about the seed.
""")

# ---------------------------------------------------------------- slide 4 (badge 3)
rewrite(shape(S[3], 7), [
    "15-bit shift register: repeats every 32,767 calls",
    "RND(n) returns the register MOD n (integers only)",
    "Plenty for games. Not for statistics.",
])
rewrite(shape(S[3], 9), ["Seeded by your keypress timing"])
rewrite(shape(S[3], 10), [
    "There's a counter at $4E and $4F that goes up while the monitor sits waiting for a keypress. "
    "Integer BASIC's generator is that counter: RND reads it, scrambles it, writes it back."
])
rewrite(shape(S[3], 12), [
    "Sander-Cederlof, Apple Assembly Line, August 1981  ·  Integer BASIC disassembly: Paul Santa-Maria, via Andy McFadden"
])
notes(S[3], """
Say out loud that $4E/$4F is not a jiffy clock. That's a Commodore term. The Apple II has no timer doing this.

Speaker line: "This isn't a toy. It's a 15-bit maximal shift register with a zero-lock guard, in 50 hand-assembled bytes." The whole Integer BASIC RNG is 56 bytes: 6 in KEYIN, 50 in RND, in an 8K ROM that couldn't even fit the hi-res routines.

Concede it before someone else does: every Integer BASIC program walks the same 32,767-step cycle; the keypress only picks where you start. Good for games, not for statistics.

Don't make the regression argument yet. Just show what worked.
""")

# ---------------------------------------------------------------- slide 5 (badge 4)
hdr = shape(S[4], 15)
rewrite(hdr, ["Same counter in the II, II+, IIe and IIc ROMs shown"])
hdr.width = Inches(5.8)
notes(S[4], """
Walk the left side first. Point at the bne going back to KEYIN. That's the whole entropy source.

The counter gets bumped before the keyboard is read, every pass. It's a loop count, not a keystroke count. The loop is 15 CPU cycles, so about 68,000 counts a second; the 16-bit counter wraps roughly once a second.

It only moves while a KEYIN-style routine waits for a key. A game that polls $C000 directly never stirs it.

Then the right side, in one line: the second thing Integer BASIC's RND does is read $4E. Same counter in the Autostart ROM, the IIe and the IIc; links are on the slide. Don't walk each image.

Sources: 6502disassembly.com OrigF8ROM and IntegerBASIC. The monitor listing credits S. Wozniak and A. Baum. Woz wrote Integer BASIC with no assembler, so there is no original source, only hand-written pages. The disassembly is Paul Santa-Maria's work, converted by Andy McFadden.

Only claim the ROMs shown. If asked about the IIgs or later IIc ROMs: not checked.
""")

# ---------------------------------------------------------------- slide 6 (badge 5)
rewrite(shape(S[5], "Text 2"), ["Applesoft II: Microsoft's RND"])
sub = shape(S[5], 5)
rewrite(sub, ["Microsoft wrote this generator. Apple licensed it and shipped it unchanged in ROM."])
for r in sub.text_frame.paragraphs[0].runs:
    r.font.size, r.font.name = Pt(15), "Arial"
    r.font.color.rgb = RGBColor.from_string(MUTED)
rewrite(shape(S[5], 8), ["Seed: $C9, filled from ROM"])
rewrite(shape(S[5], 9), ["$4E/$4F still counts. Applesoft never reads it."])
rewrite(shape(S[5], 12), [
    "Loads the seed at $C9, multiplies, adds (lost to precision), swaps the high and low "
    "mantissa bytes, forces the result under 1, writes it back to $C9. Nothing else touches $C9."
])
add_text(S[5], 0.7, 6.5, 11.9, 0.45, [
    "Listing and comments: Bob Sander-Cederlof, S-C DocuMentor, via Andy McFadden (6502disassembly.com). "
    "The comments are Sander-Cederlof's, not Microsoft's."
], 11, MUTED)
pic = shape(S[5], 6)
pic._element.nvPicPr.cNvPr.set("descr", "Applesoft RND routine, $EFAE-$EFE7, from McFadden's disassembly of S-C DocuMentor")
notes(S[5], """
This is the whole thing. Twenty-eight instructions.

Point at the top: the seed comes from $C9, which got four of its five bytes from ROM at power-on. Point at the bottom: the answer goes back to $C9. Nothing else feeds it.

Read the comments aloud: "very poor RND algorithm", "this does nothing, due to small exponent". Say whose they are: Bob Sander-Cederlof's annotations, not Microsoft's source. (The add isn't literally nothing: emulated, it changed 5 of 57,021 steps. "Lost to precision" is safe.)

This is Applesoft II as it sits in the ][+ ROM (1979) and the IIe ROMs, byte-identical. If asked about the 1978 cassette version: not checked.

Microsoft's source had two RNDs: the Commodore build read hardware timers for RND(0); every other target, including Apple, got this fixed-seed one. Apple edited the very loop that loads the seed and never connected the counter. Say "nobody connected it", not "nobody knew".

If someone wants the source: McFadden's disassembly at 6502disassembly.com, and Microsoft's own BASIC-M6502 source release.
""")

# ---------------------------------------------------------------- slide 7 (badge 5 -> 6)
rewrite(shape(S[6], 3), ["6"])
add_text(S[6], 9.44, 3.08, 3.3, 0.4, ["…starting here, every power-on."], 14, GREEN, bold=True)
add_panel(S[6], 9.44, 4.1, 3.2, 1.75)
add_text(S[6], 9.6, 4.25, 2.95, 1.5,
         ["Repeatable on purpose: documented.", "", "Repeatable by accident: not."], 17, INK, bold=True)
notes(S[6], """
Read the RND(n) line aloud: "generates a new random number each time it is used." True. It just never tells you the sequence starts in the same place every power-on.

RND(-n): Apple documents repeatability as a feature, for debugging.

The manual tells you how to get the same sequence on purpose. It never tells you you're getting it by accident. Aldridge's word for this, in 1987: "undocumented."

When Aldridge's lab called Apple, the advice was X = RND(-1*(PEEK(78)+256*PEEK(79))). Hold that for the fix slide.

(Consider showing this slide before the code slide: the promise, then the reality.)
""")

# ---------------------------------------------------------------- slide 8 (badge 7)
rewrite(shape(S[7], "Text 2"), ["Found. Published. Worked around."])
rows = [
    (5, 6, "Kaner & Vokey (1982; MICRO, 1984)", "RND fell into an endless loop of 202 numbers."),
    (7, 8, "Call-A.P.P.L.E., Jan 1983, pp. 29–34", "“RND is Fatally Flawed.”"),
    (9, 10, "Sander-Cederlof, AAL May 1984", "Seed copy is off by one ($F151: $1C→$1D). Repeats at number 37,758."),
    (11, 12, "Aldridge, BRMIC 19(4), 1987", "Each day's first subject got the same “random” word list."),
    (13, 14, "Gleason, Collegiate Microcomputer, 1988", "Tested Apple IIe RND; published suggested seeds. ERIC EJ372427."),
]
for left, right, a, b in rows:
    rewrite(shape(S[7], left), [a])
    rewrite(shape(S[7], right), [b])
tmpl_l, tmpl_r = shape(S[7], 13), shape(S[7], 14)
for src, y in ((tmpl_l, 5.70), (tmpl_r, 5.70)):
    el = copy.deepcopy(src._element)
    S[7].shapes._spTree.append(el)
new_l, new_r = list(S[7].shapes)[-2:]
new_l.top = new_r.top = Inches(5.70)
new_l.shape_id, new_r.shape_id  # ids are renumbered below
rewrite(new_l, ["Empson, GS WorldView, 1999"])
rewrite(new_r, ["Wrote up R. C. Moore's 1989 shift-register RNG, seeded from $4E/$4F."])
notes(S[7], """
Lead with the harm: Aldridge's memory experiment. Every subject was supposed to get an individually randomized word list. The ones who got the repeated list were the first tested each day, right after the computer was turned on. A published experiment, contaminated by this bug.

Then the timeline: found, reported, worked around one program at a time, never fixed in the ROM.

202 is measured on real Applesoft by two psychologists who needed randomized experiments. Sander-Cederlof got 37,758 two years later. Same generator, different uncopied byte: emulating the ROM code, $CD=$58 falls into the 202-number loop and $CD=$FF into the 37,758 one.

Say the personal line out loud: Dr. Kaner sent you the paper directly. The paper is marked 1982 but was published in MICRO in June 1984, so don't call it the first published account; Call-A.P.P.L.E. ran in January 1983.
""")

# ---------------------------------------------------------------- slide 9 (badge 9 -> 8)
rewrite(shape(S[8], 3), ["8"])
rewrite(shape(S[8], "Text 2"), ["The Fix: Put the Counter Back"])
rewrite(shape(S[8], 20), ["Part 1 · LC_Loader.bin"])
rewrite(shape(S[8], 19), [
    "Enables language card write mode while keeping ROM readable",
    "Copies $D000-$FFFF (all of Applesoft and the monitor ROM) from ROM to LC RAM",
    "Writes the patch code into the LC at $F5CB, over HFIND, which no Applesoft routine calls",
    "Writes JMP $F5CB at $EFAE in the LC copy",
    "Switches the language card to read mode",
])
rewrite(shape(S[8], 21), ["Part 2 · Patch_lc.bin"])
p2 = shape(S[8], 24)
rewrite(p2, [
    "RND(−n) and RND(0) behave as documented; RND with a positive argument draws on the KEYIN counter.",
    "Fixes the seed, not the generator. Needs a language card.",
])
p2.height = Inches(1.25)
add_panel(S[8], 7.6, 1.83, 4.95, 2.35)
add_text(S[8], 7.8, 1.95, 4.6, 0.4, ["“Just seed it first.”"], 16, GREEN, bold=True)
add_text(S[8], 7.8, 2.4, 4.6, 1.7, [
    "New code: X = RND(-(PEEK(78)+256*PEEK(79))) after a keypress. Apple's own advice (Aldridge 1987).",
    "Existing code: you can't edit it, and it fails silently.",
], 13, INK)
add_panel(S[8], 7.6, 4.35, 4.95, 2.0)
add_text(S[8], 7.8, 4.47, 4.6, 0.4, ["Patch budget"], 16, GREEN, bold=True)
add_text(S[8], 7.8, 4.92, 4.6, 1.4, [
    "HFIND spans $F5CB–$F5FF: 53 bytes.",
    "$F600 is an RTS that HLIN (HPLOT TO) branches to: leave it.",
    "No JSR or JMP to $F5CB anywhere in the ][+ ROM.",
], 13, INK)
notes(S[8], """
Somebody will say just seed it first. It's on the screen: say it yourself before they do, and agree with it for new code. Apple itself gave that one-liner to Aldridge's lab.

Then: you can't edit software you didn't write, the failure is silent so only people who already know can fix it, and this machine used to do it automatically.

HFIND: "not called by any Applesoft routine" (S-C DocuMentor). Only a program that calls 62923 ($F5CB) directly would break. The patch must fit $F5CB-$F5FF, 53 bytes; $F600 is the RTS that HLIN branches to.

CONFIRM BEFORE THE SHOW (see 01_Edited_Deck/CHANGES.md):
- Exactly what positive RND does: reseed on every call, once, or when the counter changed? A tight 52-card loop with no keypress must not return near-identical values.
- Patch length <= 53 bytes; HPLOT TO, DRAW, XDRAW still work after install.
- RND(-1) then five RND(1) repeats, with and without a GET in between.
- Behavior under DOS 3.3 (FP/INT), ProDOS, and after Ctrl-RESET on a IIe/IIc.

Limits to say out loud if asked: it fixes the seed, not the generator, so the 202-number loop is still reachable; a turnkey program that calls RND before any keypress gets leftover counter bits.
""")

# ---------------------------------------------------------------- slide 10 (add badge 9)
badge_shape = copy.deepcopy(shape(S[8], 2)._element)
badge_text = copy.deepcopy(shape(S[8], 3)._element)
S[9].shapes._spTree.append(badge_shape)
S[9].shapes._spTree.append(badge_text)
rewrite(list(S[9].shapes)[-1], ["9"])
add_text(S[9], 0.9, 3.8, 11.5, 1.4, [
    "1977: Integer BASIC's RND was seeded by you.",
    "1978: Applesoft's RND was seeded by ROM, and the counter went unread.",
    "It hid for years: rerun looked random; only power-on repeated.",
], 18, INK)
add_text(S[9], 0.9, 5.25, 11.5, 0.45, ["[ presenter: add link text and QR code here ]"], 15, AMBER)
notes(S[9], """
Three lines, then the link and QR. Put the URL on screen as text too: the back rows can't scan a QR, and the PDF needs a clickable link.

"Then and Now", the now: the same shape keeps recurring. Debian's OpenSSL (CVE-2008-0166): a 2006 change removed nearly all the entropy mixing, and predictable keys were generated for about two years before anyone noticed. A parallel, not a lineage.

Questions to expect:
- Does the C64 have this too? Same Microsoft generator, same constants. But Microsoft's own Commodore build made RND(0) read hardware timers and added TI, back in the first PET ROM in 1977; Commodore carried it to the C64. On a Commodore the idiom is RND(-TI); on the Apple it's RND(-PEEK(78)-256*PEEK(79)).
- Why not replace the generator? Speed, no space, and RND(-n) has to keep repeating.
- What's at $F5CB? HFIND. Answer it straight: no Applesoft routine calls it.
- Does the patch fix the 202 loop? No. It fixes the seed.
- ProDOS lives in the language card. Answer only with what you've tested.
""")

# Unique shape ids after deep copies
for s in prs.slides:
    used = set()
    nxt = max(int(e.get("id")) for e in s.shapes._spTree.iter(qn("p:cNvPr"))) + 1
    for e in s.shapes._spTree.iter(qn("p:cNvPr")):
        if e.get("id") in used:
            e.set("id", str(nxt))
            nxt += 1
        used.add(e.get("id"))

DST.parent.mkdir(parents=True, exist_ok=True)
prs.save(DST)
print("wrote", DST)
