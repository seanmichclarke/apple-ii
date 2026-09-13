"""Build the expanded companion deck from one slide list, in four formats.

  02_New_Deck/Applesoft_Loaded_Dice_Expanded.pptx   native PowerPoint/Keynote-importable
  02_New_Deck/Applesoft_Loaded_Dice_Expanded.html   self-contained, arrow keys, N = notes
  02_New_Deck/Applesoft_Loaded_Dice_Expanded.marp.md  Marp source (images from ../assets)
  02_New_Deck/KEYNOTE_OUTLINE.md                   per-slide outline + notes + sources

Every number on these slides is traced in 03_Collateral/FACT_SHEET.md.
"""
import base64
import html
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "assets" / "deck_images"
OUT = ROOT / "02_New_Deck"
BASE = "Applesoft_Loaded_Dice_Expanded"

C = dict(bg="17181A", panel="202225", rule="3A3E42", ink="E9E7E3", muted="9DA39B",
         green="5FCB7E", red="D9564A", amber="E0A73E")

RND_LISTING = """EFAE: 20 82 EB  RND  JSR SIGN        ; -1 / 0 / +1
EFB1: AA             TAX
EFB2: 30 18          BMI $EFCC       ; negative: reseed
EFB4: A9 C9          LDA #<RNDSEED
EFB6: A0 00          LDY #>RNDSEED
EFB8: 20 F9 EA       JSR LOAD_FAC    ; FAC = seed
EFBB: 8A             TXA
EFBC: F0 E7          BEQ $EFA5 (RTS) ; RND(0): old value
EFBE: A9 A6          LDA #<CON_RND_1
EFC0: A0 EF          LDY #>CON_RND_1
EFC2: 20 7F E9       JSR FMULT       ; x 11879546.40625
EFC5: A9 AA          LDA #<CON_RND_2
EFC7: A0 EF          LDY #>CON_RND_2
EFC9: 20 BE E7       JSR FADD        ; + 3.93e-8
|EFCC: A6 A1          LDX FAC+4       ; swap lowest and
EFCE: A5 9E          LDA FAC+1       ;   highest mantissa
EFD0: 85 A1          STA FAC+4       ;   bytes
EFD2: 86 9E          STX FAC+1
EFD4: A9 00          LDA #0
EFD6: 85 A2          STA FAC_SIGN    ; force positive
EFD8: A5 9D          LDA FAC
EFDA: 85 AC          STA FAC_EXT     ; exponent -> guard
EFDC: A9 80          LDA #$80
EFDE: 85 9D          STA FAC         ; value < 1
EFE0: 20 2E E8       JSR NORMALIZE
EFE3: A2 C9          LDX #<RNDSEED
EFE5: A0 00          LDY #>RNDSEED
EFE7: 4C 2B EB       JMP STORE_FAC   ; round, save seed"""

# ------------------------------------------------------------------ slide content
SLIDES = [
    dict(kind="title", title="Applesoft's Loaded Dice",
         subtitle="The Truth About Randomness on the Apple II, Then and Now",
         byline="Jeff Robison  ·  VCF Midwest 21  ·  2026",
         tag="Expanded edition · draft for presenter review",
         notes="Name and title, fifteen seconds. BEFORE this slide: power-cycle the machine (switch off, wait, on), PRINT RND(1), three times. Never Ctrl-RESET or PR#6 for the demo: they keep the fifth seed byte and the number changes. Rehearse on the machine you bring; its number may differ from .973136996, what matters is that it repeats."),
    dict(kind="hero", title="The Problem", lines=["]PRINT RND(1)", ".973136996"],
         caption="Power off. Power on. Same number.",
         footer="“exactly the same sequence each time the machine is powered on.”  — J. W. Aldridge, Behavior Research Methods, Instruments, & Computers 19(4), 1987",
         notes="The number they just watched three times. Don't explain it yet. If someone's machine prints something else: one seed byte is never initialized, so the exact value can vary by machine, but it is the same on every power-on of that machine."),
    dict(kind="twocol", title="Why nobody noticed",
         left=("Looks random", ["RUN it again: different numbers", "Ctrl-RESET, then RUN: different", "Reboot DOS: different"]),
         right=("Repeats", ["Power switch off, then on", "The first run of every day", "The one run nobody compares"]),
         footer="One seed byte survives everything except a power cycle. Aldridge (1987): repetition appears “only when a machine is powered off and back on.”",
         notes="This is the most counterintuitive fact in the talk, and it explains why the bug survived. A programmer at a desk reruns and reboots all day and sees variety. The first user of the day gets the same sequence. Ctrl-RESET preserving the seed was confirmed by emulating the ][+ ROM; the reboot behavior is Aldridge's observation."),
    dict(kind="bullets", title="Pseudo-random means repeatable",
         bullets=["A PRNG is a formula: same starting number in, same sequence out.",
                  "Two separate questions. Where does the sequence start? That's the seed. Does the output show patterns? That's the generator.",
                  "Entropy only answers the first question.",
                  "This talk is mostly about the seed. The generator gets its turn on slide 9."],
         footer="“Good enough” depends on the job: a game shuffle, a psychology experiment and cryptography need different things.",
         notes="Thirty seconds. The room knows what a PRNG is; this slide exists to separate seed from generator so nobody can say the talk confused them."),
    dict(kind="image", title="Integer BASIC, 1977: Woz's RND at $EF4E", images=["image2.png"],
         bullets=["15-bit shift register: one cycle of 32,767",
                  "Its state IS $4E/$4F, the monitor's keyboard-wait counter",
                  "RND(n) returns the register MOD n",
                  "Whole RNG: 6 bytes in KEYIN + 50 in RND, hand-assembled"],
         footer="Disassembly: Paul Santa-Maria, via Andy McFadden (6502disassembly.com). Period and feedback taps checked by emulating $EF4E–$EF72.",
         notes="Speaker line: a 15-bit maximal shift register with a zero-lock guard, in 50 hand-assembled bytes, in an 8K ROM that couldn't fit the hi-res routines. Concede: every program walks the same 32,767-step cycle and MOD n is biased when n isn't a power of two. Good for games, not statistics. Don't make the regression argument yet."),
    dict(kind="image", title="The entropy source: KEYIN at $FD1B", images=["image1.png"],
         bullets=["INC RNDL / INC RNDH on every pass while waiting for a key",
                  "15 CPU cycles per pass ≈ 68,000 counts a second",
                  "The 16-bit counter wraps about once a second",
                  "It moves only while KEYIN waits. Code that polls $C000 never stirs it."],
         footer="Not a jiffy clock: the Apple II has no timer. Monitor listing credits S. Wozniak and A. Baum (Apple II Reference Manual, 1978).",
         notes="Point at the BNE back to KEYIN: that is the entire entropy source. It's a loop count, not a keystroke count. Human reaction-time jitter spans many wraps, which is why the value looks uniform. Same counter exists in the Autostart, IIe and IIc firmware; don't claim the IIgs."),
    dict(kind="image", title="Applesoft II: the manual's promise", images=["image8.png", "image9.png"],
         bullets=["RND(n): “a new random number each time it is used”",
                  "RND(−n): a repeatable sequence, on purpose, for debugging",
                  "Never mentioned: every power-on starts the same sequence"],
         callout="Repeatable on purpose: documented. Repeatable by accident: not.",
         footer="Apple, Applesoft II BASIC Programming Reference Manual (scans from the presenter's deck; confirm edition before citing a year).",
         notes="The promise before the reality. The manual tells you how to get the same sequence on purpose and never tells you you're getting it by accident. Aldridge's word in 1987: undocumented."),
    dict(kind="code", title="Microsoft's RND at $EFAE", code=RND_LISTING.split("|")[0].rstrip(),
         bullets=["Microsoft's 6502 BASIC, licensed by Apple, shipped unchanged",
                  "Seed at $C9–$CD. Only RND reads or writes it.",
                  "$4E/$4F still counts under Applesoft. RND never reads it.",
                  "The add is lost to precision: it changed 5 of 57,021 steps"],
         footer="Bytes checked against the ][+ ROM image. Labels after S-C DocuMentor (Sander-Cederlof) via McFadden. Full 28-instruction listing: backup slide A1.",
         notes="Twenty-eight instructions in all; this shows the first fourteen, the swap-and-normalize half is on backup A1. Point at the top: the seed comes from $C9. The answer goes back to $C9. Nothing else feeds it. The famous comments ('very poor RND algorithm') are Sander-Cederlof's, not Microsoft's."),
    dict(kind="table", title="The seed copy is off by one",
         kicker="F150: A2 1C   LDX #$1C  copies 4 of the 5 seed bytes. $CD keeps whatever RAM held.",
         columns=["$CD at power-on", "First PRINT RND(1)"],
         rows=[["$FF or $FE", ".973136996"], ["$00", ".270011996"], ["$58 (the byte ROM meant to copy)", ".512199496"], ["$AA", ".738761996"]],
         highlight=0,
         footer="256 possible $CD values give 181 different first numbers. Emulated on the ][+ ROM (SHA-1 33a24f54…). Fix published in AAL, May 1984: $F151 from $1C to $1D.",
         notes="This is why the audience's own machine may print a different number, and why only a power cycle repeats. The bug is in Microsoft's source (LDXI RNDX+4-CHRGET), and Steil reports it in every Microsoft 6502 BASIC. Apple edited this very loop (it inserted STX SPEEDZ) and kept the short count."),
    dict(kind="chart", title="The generator falls into short loops",
         kicker="Every cold start ends in one of five loops. The seed has about 4 billion possible values.",
         data=[("37,758", 43.0), ("32,366", 30.5), ("202", 23.0), ("4,082", 2.0), ("12,559", 1.6)],
         axis="Share of the 256 possible $CD values that end in each loop (loop length in calls)",
         footer="Kaner & Vokey saw 202; Sander-Cederlof saw 37,758. Same generator, different uncopied byte. Exhaustive emulated sweep of $CD; other reseeds may reach other loops.",
         notes="The two published periods don't disagree; they're two of the five loops. Almost a quarter of possible power-on bytes put an untouched program into a loop of 202 numbers. A better seed doesn't escape this: RND(-52894) and RND(-22258) also land in the 202 loop."),
    dict(kind="quote", title="Who got hurt",
         quote="…the subjects receiving the repeated lists were those tested at the beginning of each day, immediately after the computer had been turned on.",
         attribution="J. W. Aldridge, “Cautions regarding random number generation on the Apple II,” Behavior Research Methods, Instruments, & Computers 19(4):397–399, 1987",
         footer="A memory experiment meant to give every subject an individually randomized word list.",
         notes="This is the stakes. Not a game with a predictable shuffle, which we'd be guessing at, but a published experiment. Kaner and Vokey were writing for the same reason: every experiment they ran needed randomized order."),
    dict(kind="table", title="A regression of the default, not the capability",
         columns=["", "Integer BASIC (1977)", "Applesoft II (1978–)"],
         rows=[["Generator", "15-bit shift register", "Microsoft floating-point multiply, swap, normalize"],
               ["State", "15 bits", "5-byte float (32-bit mantissa)"],
               ["Cycle", "32,767, one cycle", "Loops of 202 to 37,758"],
               ["Power-on seed", "Whatever $4E/$4F holds", "4 ROM bytes + 1 leftover byte"],
               ["Keyboard timing", "Used automatically", "Only if the program asks"],
               ["Weakness", "Tiny state, MOD bias", "Short loops, fixed start"]],
         footer="Both generators are weak. Only one threw away the seed the machine was already collecting.",
         notes="This is where the talk survives 'both are bad'. Concede the Integer BASIC weaknesses openly. The claim is narrow: Integer BASIC gave you keystroke timing by default because the counter and the generator state are the same bytes; Applesoft kept the counter running and ignored it."),
    dict(kind="bullets", title="Why 1978 got worse",
         bullets=["Microsoft's source had two RNDs: the Commodore build read hardware timers; every other target got a fixed seed.",
                  "Apple got the fixed-seed build.",
                  "Apple edited the seed-copy loop itself (it inserted STX SPEEDZ) and kept the short count.",
                  "RND, its constants and that loop are byte-identical in the ][+, IIe and enhanced IIe ROMs."],
         footer="No record says why. Say “nobody connected it,” not “nobody knew.” Sources: Microsoft BASIC-M6502 source (2025 release); AppleWin ROM images compared byte-wise.",
         notes="Portability doesn't hold up as a reason: Microsoft's file already had Apple-only code, and Commodore got a machine-specific RND in the same file. Motive is inference; label it that way on stage."),
    dict(kind="twocol", title="“Just seed it first.”",
         left=("New code", ["X = RND(-(PEEK(78)+256*PEEK(79)))", "Apple's own advice to Aldridge's lab", "Only works after a keypress", "Both bytes 0 means RND(0): no reseed"]),
         right=("Existing code", ["Thousands of programs never did it", "You can't edit software you didn't write", "The failure is silent", "The machine used to do it for you"]),
         footer="Agree with the heckle for new code. The patch is for everything already written.",
         notes="Say this before anyone in the room does. The zero case is a genuine edge: SIGN reduces 0 to the RND(0) path, which returns the old value and never reseeds."),
    dict(kind="table", title="The fix: put the counter back",
         kicker="Part 1, LC_Loader.bin copies ROM into language-card RAM. Part 2, Patch_lc.bin is the new RND front end.",
         columns=["Address", "In ROM", "In the language-card copy"],
         rows=[["$D000–$EFAD", "Applesoft", "copied"],
               ["$EFAE", "RND: JSR SIGN", "JMP $F5CB"],
               ["$EFB1–$F5CA", "Applesoft", "copied"],
               ["$F5CB–$F5FF", "HFIND (53 bytes, no Applesoft caller)", "patch code"],
               ["$F600", "RTS shared with HLIN", "must stay RTS"],
               ["$F601–$FFFF", "DRAW … Monitor", "copied"]],
         highlight=[1, 3, 4],
         footer="Checked on the ][+ ROM: no JSR or JMP targets $F5CB; the branch at $F59C lands on the RTS at $F600. RND(−n) and RND(0) keep their documented behavior.",
         notes="Presenter to confirm: exact patch length (at most 53 bytes); exactly what positive RND does with the counter (every call, once, or when it changed); that HPLOT TO, DRAW and XDRAW still work after install."),
    dict(kind="bullets", title="Limits, and what to test before release",
         bullets=["Fixes the seed, not the generator: the 202-number loop is still reachable.",
                  "Needs a keypress first: a turnkey HELLO that calls RND immediately gets leftover counter bits.",
                  "Needs a language card, and shares it: test under DOS 3.3 with INT/FP, and under ProDOS.",
                  "Test what Ctrl-RESET does to the patch on a IIe and a IIc.",
                  "A tight loop with no keypress must not repeat values."],
         footer="Compatibility matrix and test scripts: 03_Collateral/HARDWARE_TEST_PLAN.md. Mark each cell tested, emulated, or untested.",
         notes="Scope, stated plainly, is a strength: Kaner & Vokey, Call-A.P.P.L.E., Sander-Cederlof and Moore all replaced the generator. This patch deliberately doesn't, so RND(-n) keeps repeating and no program changes."),
    dict(kind="hero", title="Proof", lines=["power on  →  BRUN LC_LOADER", "]PRINT RND(1)", "power off, on, again  →  a different number"],
         caption="Close the loop the opening demo opened.",
         footer="If power-cycling on stage is risky, show a recorded clip. About 90 seconds.",
         notes="Symmetry is the cheapest persuasion available: the talk opened with proof of the problem, so close with proof of the fix."),
    dict(kind="bullets", title="Now",
         bullets=["Debian OpenSSL, 2006–2008 (CVE-2008-0166): a code change removed nearly all entropy input. Keys still looked random, for about two years.",
                  "C's rand() without srand() behaves as if seeded with 1: the same sequence every run, today.",
                  "Same shape as 1978: the output looks fine; the seed quietly stopped being random."],
         footer="A parallel, not a lineage. Nobody learned this from the Apple II, and nobody ignored it either.",
         notes="Thirty seconds. This delivers the 'and Now' in the subtitle without overclaiming."),
    dict(kind="takeaways", title="Take home",
         bullets=["1977: Integer BASIC's RND was seeded by you.",
                  "1978: Applesoft's RND was seeded by ROM, and the counter went unread.",
                  "It hid for years: rerun looked random; only power-on repeated."],
         link="Try it yourself:  [ presenter: link text + QR code ]",
         byline="Jeff Robison  ·  VCF Midwest 21  ·  2026",
         notes="Three lines, then the link. Put the URL on screen as text, not only as a QR."),
    # ------------------------------------------------------------ backup
    dict(kind="code2", title="A1 · Backup: RND, all 28 instructions", code=RND_LISTING.replace("|", ""),
         footer="Bytes match the AppleWin Apple2_Plus.rom image; identical in the IIe and enhanced IIe images. Labels after S-C DocuMentor.",
         notes="For the person who asks to see every byte."),
    dict(kind="table", title="A2 · Backup: the constants as actually read",
         kicker="Each constant is 4 bytes in ROM, but the float loader always reads 5.",
         columns=["Constant", "Bytes read", "Value"],
         rows=[["Multiplier ($EFA6)", "98 35 44 7A + 68", "11,879,546.40625"],
               ["Addend ($EFAA)", "68 28 B1 46 + 20", "3.92767778 × 10⁻⁸"],
               ["Seed table ($F123)", "80 4F C7 52 58", "≈ 0.811635157 (last byte never copied)"],
               ["RND(−1) seed", "—", "2.99196472E-08"]],
         footer="Truncated constants are in Microsoft's own source (octal, RADIX 8). The stray 5th bytes are the next constant and the JSR opcode of RND.",
         notes="Sander-Cederlof: 'THESE ARE MISSING ONE BYTE FOR FP VALUES'."),
    dict(kind="table", title="A3 · Backup: every cold start, all 256 $CD values",
         columns=["Loop length", "$CD values", "Share", "Where it's been seen"],
         rows=[["37,758", "110", "43.0%", "Sander-Cederlof, AAL 1984; $CD=$FF, $00"],
               ["32,366", "78", "30.5%", "RND(-2), RND(-54321)"],
               ["202", "59", "23.0%", "Kaner & Vokey 1984; $CD=$58, $7A"],
               ["4,082", "5", "2.0%", "RND(-25284)"],
               ["12,559", "4", "1.6%", "$CD=$5F, $6D, $A8, $A9"]],
         footer="Emulated with py65 on the real ROM routine. Harness sets FAC and seed only; no hardware run. Raw rows: 04_Research_and_Evidence/tools/cd_sweep_results.jsonl.",
         notes="If challenged on method: the tools are in the handout folder and re-run in about ten seconds per start."),
    dict(kind="table", title="A4 · Backup: open items before the show",
         columns=["Question", "Why it matters", "Status"],
         rows=[["What $CD holds at power-on on the demo machine", "Decides the number on screen", "untested"],
               ["Patch length ≤ 53 bytes; HPLOT TO / DRAW work", "HFIND budget", "untested"],
               ["Positive RND in a tight loop, no keypress", "Could repeat values", "untested"],
               ["RND(−1) + five RND(1), with and without GET", "Documented contract", "untested"],
               ["DOS 3.3 INT/FP, ProDOS, Ctrl-RESET on IIe/IIc", "Language-card conflicts", "untested"],
               ["//c and IIgs Applesoft bytes", "“All Apple II” claims", "not checked"]],
         footer="Hide or delete this slide before presenting.",
         notes="Working slide for the presenter, not the audience."),
    dict(kind="bullets", title="A5 · Sources",
         bullets=["Kaner & Vokey, “A Better Random Number Generator for Apple's Floating Point BASIC,” MICRO, June 1984 (ms. 1982)",
                  "“RND is Fatally Flawed,” Call-A.P.P.L.E., Jan 1983, pp. 29–34",
                  "B. Sander-Cederlof, “Random Numbers for Applesoft,” Apple Assembly Line, May 1984; S-C DocuMentor",
                  "J. W. Aldridge, Behavior Research Methods, Instruments, & Computers 19(4):397–399, 1987",
                  "J. M. Gleason, Collegiate Microcomputer 6(2):108–112, 1988 (ERIC EJ372427)",
                  "Microsoft, BASIC-M6502 source, github.com/microsoft/BASIC-M6502",
                  "A. McFadden, 6502disassembly.com (Applesoft, Integer BASIC, monitor ROMs)",
                  "M. Steil, pagetable.com; D. Empson, GS WorldView, Nov 1999"],
         footer="Full annotated bibliography: 03_Collateral/SOURCES.md",
         notes=""),
]

W, H = 13.333, 7.5


# ------------------------------------------------------------------ PPTX
def rgb(k):
    return RGBColor.from_string(C[k])


def tb(slide, x, y, w, h, paras, size=18, color="ink", bold=False, mono=False, align=None, anchor=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    for i, t in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        r = p.add_run()
        r.text = t
        f = r.font
        f.size, f.bold = Pt(size), bold
        f.name = "Courier New" if mono else "Arial"
        f.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, fill="panel"):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    return s


def bullets_box(slide, x, y, w, h, items, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = "▸  "
        r1.font.color.rgb, r1.font.size, r1.font.name = rgb("green"), Pt(size), "Arial"
        r2 = p.add_run()
        r2.text = t
        r2.font.color.rgb, r2.font.size, r2.font.name = rgb("ink"), Pt(size), "Arial"


def header(slide, s, n):
    tb(slide, 0.7, 0.45, 11.0, 0.8, s["title"], 32, "ink", bold=True)
    rect(slide, 11.95, 0.42, 0.62, 0.62)
    tb(slide, 11.95, 0.42, 0.62, 0.62, str(n), 15, "green", bold=True, mono=True,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if s.get("kicker"):
        tb(slide, 0.7, 1.25, 11.9, 0.6, s["kicker"], 16, "muted")


def footer(slide, s):
    if s.get("footer"):
        tb(slide, 0.7, 6.55, 11.9, 0.8, s["footer"], 12, "muted")


def img_fit(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(dw), Inches(dh))
    return dh


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for n, s in enumerate(SLIDES):
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb("bg")
        k = s["kind"]
        if k == "title":
            tb(sl, 0.9, 2.3, 11.5, 1.1, s["title"], 52, "ink", bold=True)
            tb(sl, 0.9, 3.4, 11.5, 0.6, s["subtitle"], 22, "green")
            rect(sl, 0.92, 4.25, 1.6, 0.03, "rule")
            tb(sl, 0.9, 4.5, 11.5, 0.5, s["byline"], 16, "muted")
            tb(sl, 0.9, 6.6, 11.5, 0.4, s["tag"], 12, "amber")
        elif k == "takeaways":
            tb(sl, 0.9, 1.2, 11.5, 0.9, s["title"], 40, "ink", bold=True)
            rect(sl, 0.92, 2.15, 1.6, 0.03, "green")
            bullets_box(sl, 0.9, 2.5, 11.5, 2.6, s["bullets"], 24)
            tb(sl, 0.9, 5.2, 11.5, 0.5, s["link"], 18, "amber")
            tb(sl, 0.9, 6.3, 11.5, 0.4, s["byline"], 15, "muted")
        else:
            header(sl, s, n)
            top = 1.95 if s.get("kicker") else 1.6
            if k == "hero":
                rect(sl, 3.2, 2.1, 6.9, 2.6)
                tb(sl, 3.5, 2.3, 6.3, 2.2, s["lines"], 26 if len(s["lines"]) < 3 else 20, "green", mono=True,
                   anchor=MSO_ANCHOR.MIDDLE)
                tb(sl, 0.7, 5.0, 11.9, 0.6, s["caption"], 22, "ink", bold=True, align=PP_ALIGN.CENTER)
            elif k == "bullets":
                bullets_box(sl, 0.7, top, 11.9, 4.6, s["bullets"], 20)
            elif k == "twocol":
                for i, (head, items) in enumerate([s["left"], s["right"]]):
                    x = 0.7 + i * 6.1
                    rect(sl, x, top, 5.8, 4.1)
                    tb(sl, x + 0.3, top + 0.2, 5.2, 0.5, head, 20, "green", bold=True)
                    bullets_box(sl, x + 0.3, top + 0.85, 5.3, 3.1, items, 18)
            elif k == "image":
                y = top
                for im in s["images"]:
                    dh = img_fit(sl, IMG / im, 0.7, y, 6.6, 4.3 / len(s["images"]) - 0.1)
                    y += dh + 0.2
                rect(sl, 7.6, top, 5.0, 4.4 if not s.get("callout") else 3.2)
                bullets_box(sl, 7.8, top + 0.2, 4.6, 4.0, s["bullets"], 16)
                if s.get("callout"):
                    rect(sl, 7.6, top + 3.35, 5.0, 1.05, "rule")
                    tb(sl, 7.8, top + 3.42, 4.6, 0.9, s["callout"], 16, "ink", bold=True, anchor=MSO_ANCHOR.MIDDLE)
            elif k in ("code", "code2"):
                lines = s["code"].splitlines()
                if k == "code":
                    rect(sl, 0.7, top, 7.2, 4.75)
                    tb(sl, 0.85, top + 0.1, 7.0, 4.6, lines, 12, "green", mono=True)
                    bullets_box(sl, 8.15, top, 4.5, 4.7, s["bullets"], 16)
                else:
                    half = (len(lines) + 1) // 2
                    for i, chunk in enumerate([lines[:half], lines[half:]]):
                        rect(sl, 0.7 + i * 6.05, top, 5.85, 4.75)
                        tb(sl, 0.8 + i * 6.05, top + 0.1, 5.7, 4.6, chunk, 11, "green", mono=True)
            elif k == "table":
                rows, cols = s["rows"], s["columns"]
                hl = s.get("highlight")
                hl = [hl] if isinstance(hl, int) else (hl or [])
                rh = min(0.62, 4.4 / (len(rows) + 1))
                shp = sl.shapes.add_table(len(rows) + 1, len(cols), Inches(0.7), Inches(top),
                                          Inches(11.9), Inches(rh * (len(rows) + 1)))
                t = shp.table
                fs = 15 if len(rows) <= 5 else 13
                weights = [max(len(c), *(len(r[i]) for r in rows)) + 6 for i, c in enumerate(cols)]
                for ci, wgt in enumerate(weights):
                    t.columns[ci].width = Inches(11.9 * wgt / sum(weights))
                for ci, head in enumerate(cols):
                    cell = t.cell(0, ci)
                    cell.fill.solid(); cell.fill.fore_color.rgb = rgb("rule")
                    cell.text = head
                    f = cell.text_frame.paragraphs[0].runs[0].font if head else None
                    if f:
                        f.size, f.bold, f.name = Pt(15), True, "Arial"; f.color.rgb = rgb("ink")
                for ri, row in enumerate(rows, 1):
                    for ci, val in enumerate(row):
                        cell = t.cell(ri, ci)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb("panel") if (ri - 1) not in hl else RGBColor(0x24, 0x33, 0x2A)
                        cell.text = val
                        r = cell.text_frame.paragraphs[0].runs[0]
                        mono = val.startswith(("$", ".", "9", "6", "8")) or val[:1].isdigit()
                        r.font.size, r.font.name = Pt(fs), ("Courier New" if mono else "Arial")
                        r.font.color.rgb = rgb("green") if (ri - 1) in hl else rgb("ink")
            elif k == "chart":
                cd = CategoryChartData()
                cd.categories = [d[0] for d in s["data"]][::-1]
                cd.add_series("Share of $CD values (%)", [d[1] for d in s["data"]][::-1])
                gf = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(top),
                                         Inches(11.9), Inches(4.3), cd)
                ch = gf.chart
                ch.has_legend = False
                ch.font.size, ch.font.name = Pt(14), "Arial"
                ch.font.color.rgb = rgb("muted")
                pl = ch.plots[0]
                pl.gap_width = 45
                pl.has_data_labels = True
                dl = pl.data_labels
                dl.number_format, dl.number_format_is_linked = '0.0"%"', False
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
                dl.font.size = Pt(14); dl.font.color.rgb = rgb("ink")
                ser = pl.series[0]
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb("green")
                va = ch.value_axis
                va.maximum_scale, va.minimum_scale = 50, 0
                va.has_major_gridlines = True
                va.major_gridlines.format.line.color.rgb = rgb("rule")
                va.format.line.fill.background()
                va.tick_labels.font.color.rgb = rgb("muted")
                ca = ch.category_axis
                ca.format.line.color.rgb = rgb("rule")
                ca.tick_labels.font.color.rgb = rgb("ink"); ca.tick_labels.font.name = "Courier New"
                tb(sl, 0.7, top + 4.3, 11.9, 0.4, s["axis"], 13, "muted")
            elif k == "quote":
                rect(sl, 0.7, 1.9, 0.08, 2.6, "green")
                tb(sl, 1.1, 1.85, 11.3, 2.7, "“" + s["quote"] + "”", 28, "ink")
                tb(sl, 1.1, 4.7, 11.3, 0.9, s["attribution"], 15, "muted")
            footer(sl, s)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.core_properties.title = "Applesoft's Loaded Dice - expanded edition"
    prs.core_properties.author = "Draft for Jeff Robison"
    path = OUT / f"{BASE}.pptx"
    prs.save(path)
    return path


# ------------------------------------------------------------------ HTML
def b64(name):
    return "data:image/png;base64," + base64.b64encode((IMG / name).read_bytes()).decode()


def e(t):
    return html.escape(t)


def html_slide(n, s):
    k = s["kind"]
    body = ""
    if k == "title":
        return (f'<section class="slide title"><h1>{e(s["title"])}</h1><p class="sub">{e(s["subtitle"])}</p>'
                f'<hr><p class="by">{e(s["byline"])}</p><p class="tag">{e(s["tag"])}</p>'
                f'<aside class="notes">{e(s["notes"])}</aside></section>')
    if k == "takeaways":
        lis = "".join(f"<li>{e(b)}</li>" for b in s["bullets"])
        return (f'<section class="slide take"><h1>{e(s["title"])}</h1><hr class="g"><ul>{lis}</ul>'
                f'<p class="link">{e(s["link"])}</p><p class="by">{e(s["byline"])}</p>'
                f'<aside class="notes">{e(s["notes"])}</aside></section>')
    ul = lambda items: "<ul>" + "".join(f"<li>{e(b)}</li>" for b in items) + "</ul>"
    if k == "hero":
        body = f'<div class="hero"><pre>{e(chr(10).join(s["lines"]))}</pre></div><p class="caption">{e(s["caption"])}</p>'
    elif k == "bullets":
        body = ul(s["bullets"])
    elif k == "twocol":
        body = '<div class="cols">' + "".join(
            f'<div class="panel"><h3>{e(h)}</h3>{ul(it)}</div>' for h, it in (s["left"], s["right"])) + "</div>"
    elif k == "image":
        imgs = "".join(f'<img src="{b64(i)}" alt="{e(s["title"])}">' for i in s["images"])
        call = f'<div class="callout">{e(s["callout"])}</div>' if s.get("callout") else ""
        body = f'<div class="split"><div class="imgs">{imgs}</div><div><div class="panel">{ul(s["bullets"])}</div>{call}</div></div>'
    elif k == "code":
        body = f'<div class="split"><pre class="code">{e(s["code"])}</pre><div>{ul(s["bullets"])}</div></div>'
    elif k == "code2":
        lines = s["code"].splitlines(); half = (len(lines) + 1) // 2
        body = '<div class="cols">' + "".join(f'<pre class="code">{e(chr(10).join(c))}</pre>' for c in (lines[:half], lines[half:])) + "</div>"
    elif k == "table":
        hl = s.get("highlight"); hl = [hl] if isinstance(hl, int) else (hl or [])
        th = "".join(f"<th>{e(c)}</th>" for c in s["columns"])
        trs = "".join('<tr class="%s">' % ("hl" if i in hl else "") + "".join(f"<td>{e(v)}</td>" for v in r) + "</tr>"
                      for i, r in enumerate(s["rows"]))
        body = f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>"
    elif k == "chart":
        rows = s["data"]; bw = 760; y = 0; parts = []
        for lab, v in rows:
            w = bw * v / 50
            parts.append(f'<g class="bar" tabindex="0"><title>Loop of {lab} calls: {v}% of $CD values</title>'
                         f'<rect x="0" y="{y-4}" width="{bw+220}" height="52" fill="transparent"/>'
                         f'<text x="-14" y="{y+28}" text-anchor="end" class="cat">{lab}</text>'
                         f'<path d="M0 {y} H{max(w-4,0)} a4 4 0 0 1 4 4 V{y+40} a4 4 0 0 1 -4 4 H0 Z" fill="#{C["green"]}"/>'
                         f'<text x="{w+10}" y="{y+29}" class="val">{v:.1f}%</text></g>')
            y += 60
        grid = "".join(f'<line x1="{bw*g/50}" y1="-8" x2="{bw*g/50}" y2="{y-10}" class="grid"/><text x="{bw*g/50}" y="{y+14}" class="tick" text-anchor="middle">{g}%</text>' for g in (0, 10, 20, 30, 40, 50))
        tbl = "".join(f"<tr><td>{lab}</td><td>{v:.1f}%</td></tr>" for lab, v in rows)
        body = (f'<svg class="chart" viewBox="-110 -20 1100 {y+40}" role="img" aria-label="{e(s["axis"])}">{grid}{"".join(parts)}</svg>'
                f'<p class="axis">{e(s["axis"])}</p><details class="tv"><summary>Table view</summary><table><tr><th>Loop length</th><th>Share</th></tr>{tbl}</table></details>')
    elif k == "quote":
        body = f'<blockquote>“{e(s["quote"])}”</blockquote><p class="attr">{e(s["attribution"])}</p>'
    kick = f'<p class="kicker">{e(s["kicker"])}</p>' if s.get("kicker") else ""
    foot = f'<footer>{e(s["footer"])}</footer>' if s.get("footer") else ""
    return (f'<section class="slide {k}"><header><h2>{e(s["title"])}</h2><span class="badge">{n}</span></header>'
            f'{kick}<div class="body">{body}</div>{foot}<aside class="notes">{e(s.get("notes",""))}</aside></section>')


CSS = """
:root{--bg:#%(bg)s;--panel:#%(panel)s;--rule:#%(rule)s;--ink:#%(ink)s;--muted:#%(muted)s;--green:#%(green)s;--amber:#%(amber)s}
*{box-sizing:border-box}html,body{margin:0;height:100%%;background:#000;font-family:Arial,Helvetica,sans-serif}
#stage{position:fixed;inset:0;display:flex;align-items:center;justify-content:center}
.slide{display:none;position:absolute;width:1280px;height:720px;background:var(--bg);color:var(--ink);padding:40px 64px;transform-origin:center}
.slide.on{display:block}
header{display:flex;justify-content:space-between;align-items:flex-start}
h2{font-size:34px;margin:4px 0 0}.badge{font:bold 16px 'Courier New',monospace;color:var(--green);background:var(--panel);width:58px;height:58px;display:grid;place-items:center}
.kicker{color:var(--muted);font-size:17px;margin:10px 0 0}
.body{margin-top:22px}ul{margin:0;padding:0;list-style:none}li{font-size:22px;line-height:1.35;margin:0 0 14px;padding-left:30px;position:relative}
li:before{content:'\\25B8';color:var(--green);position:absolute;left:0}
footer{position:absolute;left:64px;right:64px;bottom:26px;color:var(--muted);font-size:13.5px;line-height:1.4}
.title,.take{padding:190px 90px}.title h1{font-size:56px;margin:0}.sub{color:var(--green);font-size:24px;margin:14px 0 26px}
hr{width:150px;border:0;border-top:3px solid var(--rule);margin:0 0 22px}hr.g{border-color:var(--green)}
.by{color:var(--muted);font-size:17px}.tag,.link{color:var(--amber)}.tag{position:absolute;bottom:40px;font-size:13px}.link{font-size:20px;margin-top:26px}
.take{padding-top:100px}.take h1{font-size:44px;margin:0 0 20px}.take li{font-size:26px}
.hero{background:var(--panel);width:660px;margin:30px auto 0;padding:34px 40px}.hero pre{margin:0;color:var(--green);font:28px/1.5 'Courier New',monospace;white-space:pre-wrap}
.caption{text-align:center;font-size:24px;font-weight:bold;margin-top:34px}
.cols{display:grid;grid-template-columns:1fr 1fr;gap:24px}.panel{background:var(--panel);padding:22px 26px}.panel h3{color:var(--green);margin:0 0 18px;font-size:22px}
.panel li{font-size:19px}.split{display:grid;grid-template-columns:1.3fr 1fr;gap:28px}.imgs img{width:100%%;display:block;margin-bottom:12px;background:#fff}
.split .panel li{font-size:17px}.callout{background:var(--rule);margin-top:14px;padding:16px 22px;font-weight:bold;font-size:17px}
pre.code{background:var(--panel);color:var(--green);font:13.5px/1.42 'Courier New',monospace;padding:14px 18px;margin:0;white-space:pre}
.code .split li{font-size:18px}
table{border-collapse:separate;border-spacing:0 2px;width:100%%;font-size:17px}th{background:var(--rule);text-align:left;padding:10px 14px}
td{background:var(--panel);padding:10px 14px}tr.hl td{color:var(--green);background:#24332a}td:first-child{font-family:'Courier New',monospace}
.chart{width:100%%;height:380px}.chart .cat{fill:var(--ink);font:18px 'Courier New',monospace}.chart .val{fill:var(--ink);font-size:18px}
.chart .grid{stroke:var(--rule);stroke-width:1}.chart .tick{fill:var(--muted);font-size:14px}.bar:hover path,.bar:focus path{filter:brightness(1.25)}
.axis{color:var(--muted);font-size:14px;margin:4px 0}.tv{color:var(--muted);font-size:13px}.tv table{width:auto;font-size:13px}
blockquote{border-left:6px solid var(--green);margin:40px 0 0;padding:4px 0 4px 34px;font-size:32px;line-height:1.35}.attr{color:var(--muted);font-size:16px;padding-left:40px}
aside.notes{display:none}body.shownotes aside.notes{display:block;position:fixed;left:0;right:0;bottom:0;max-height:32vh;overflow:auto;background:#111;color:#ddd;padding:14px 22px;font-size:15px;border-top:2px solid var(--amber);z-index:9}
#hud{position:fixed;right:12px;bottom:8px;color:#666;font:12px Arial;z-index:10}
@media print{#hud{display:none}.slide{display:block!important;position:relative;transform:none!important;page-break-after:always}#stage{position:static;display:block}}
""" % C

JS = """
const S=[...document.querySelectorAll('.slide')];let i=Math.max(0,Math.min(S.length-1,(parseInt(location.hash.slice(1))||0)));
function fit(){const k=Math.min(innerWidth/1280,innerHeight/720);S.forEach(s=>s.style.transform='scale('+k+')')}
function go(n){S[i].classList.remove('on');i=Math.max(0,Math.min(S.length-1,n));S[i].classList.add('on');location.hash=i;hud.textContent=(i)+' / '+(S.length-1)+'   ← → navigate · N notes'}
addEventListener('keydown',ev=>{if(['ArrowRight','PageDown',' '].includes(ev.key))go(i+1);if(['ArrowLeft','PageUp'].includes(ev.key))go(i-1);if(ev.key==='n'||ev.key==='N')document.body.classList.toggle('shownotes');if(ev.key==='Home')go(0);if(ev.key==='End')go(S.length-1)});
addEventListener('click',ev=>{if(!ev.target.closest('details,.notes'))go(ev.clientX>innerWidth/2?i+1:i-1)});
addEventListener('resize',fit);fit();go(i);
"""


def build_html():
    slides = "\n".join(html_slide(n, s) for n, s in enumerate(SLIDES))
    doc = (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Applesoft’s Loaded Dice — expanded</title>'
           f'<meta name="viewport" content="width=device-width,initial-scale=1"><style>{CSS}</style></head>'
           f'<body><div id="stage">{slides}</div><div id="hud"></div><script>{JS}</script></body></html>')
    path = OUT / f"{BASE}.html"
    path.write_text(doc, encoding="utf-8")
    return path


# ------------------------------------------------------------------ Marp + outline
def md_body(s):
    k, out = s["kind"], []
    if s.get("kicker"):
        out.append(f"*{s['kicker']}*\n")
    if k == "hero":
        out += ["```", *s["lines"], "```", f"\n**{s['caption']}**"]
    elif k in ("bullets", "takeaways"):
        out += [f"- {b}" for b in s["bullets"]]
        if s.get("link"):
            out.append(f"\n**{s['link']}**")
    elif k == "twocol":
        for h, it in (s["left"], s["right"]):
            out.append(f"\n**{h}**\n")
            out += [f"- {b}" for b in it]
    elif k == "image":
        out += [f"![w:560](../assets/deck_images/{i})" for i in s["images"]]
        out += [""] + [f"- {b}" for b in s["bullets"]]
        if s.get("callout"):
            out.append(f"\n> **{s['callout']}**")
    elif k in ("code", "code2"):
        out += ["```asm", s["code"], "```"]
        out += [f"- {b}" for b in s.get("bullets", [])]
    elif k == "table":
        out.append("| " + " | ".join(c or " " for c in s["columns"]) + " |")
        out.append("|" + "---|" * len(s["columns"]))
        out += ["| " + " | ".join(r) + " |" for r in s["rows"]]
    elif k == "chart":
        out.append("| Loop length (calls) | Share of $CD values |\n|---|---|")
        out += [f"| {a} | {b:.1f}% |" for a, b in s["data"]]
        out.append(f"\n<small>{s['axis']}</small>")
    elif k == "quote":
        out += [f"> {s['quote']}", f"\n— {s['attribution']}"]
    if s.get("footer"):
        out.append(f"\n<small>{s['footer']}</small>")
    return "\n".join(out)


def build_marp():
    parts = ["---", "marp: true", "theme: default", "paginate: true", "size: 16:9",
             "style: |", "  section { background:#17181A; color:#E9E7E3; font-family:Arial; }",
             "  h1, h2 { color:#E9E7E3; } strong { color:#5FCB7E; } small { color:#9DA39B; }",
             "  code, pre { background:#202225; color:#5FCB7E; } th { background:#3A3E42; } td { background:#202225; }",
             "---", ""]
    for n, s in enumerate(SLIDES):
        if n:
            parts.append("\n---\n")
        if s["kind"] == "title":
            parts += [f"# {s['title']}", f"### {s['subtitle']}", "", s["byline"], "", f"<small>{s['tag']}</small>"]
        else:
            parts += [f"## {s['title']}", "", md_body(s)]
        if s.get("notes"):
            parts.append(f"\n<!-- {s['notes']} -->")
    path = OUT / f"{BASE}.marp.md"
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")
    return path


def build_outline():
    out = ["# Keynote outline: Applesoft's Loaded Dice (expanded edition)", "",
           "One section per slide: on-screen content, speaker notes, and the assets it uses. "
           "Import the .pptx into Keynote for the laid-out version; use this file to rebuild by hand.",
           "Slides 0-18 are the talk (about 35 min with both demos); A1-A5 are backup.", ""]
    for n, s in enumerate(SLIDES):
        out.append(f"## {n} · {s['title']}")
        out.append(f"*Layout:* {s['kind']}" + (f" · *Assets:* {', '.join('assets/deck_images/' + i for i in s['images'])}" if s.get("images") else ""))
        out.append("")
        if s["kind"] == "title":
            out += [s["subtitle"], s["byline"]]
        else:
            out.append(md_body(s))
        if s.get("notes"):
            out += ["", f"**Speaker notes.** {s['notes']}"]
        out.append("")
    path = OUT / "KEYNOTE_OUTLINE.md"
    path.write_text("\n".join(out), encoding="utf-8")
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for f in (build_pptx, build_html, build_marp, build_outline):
        print("wrote", f())
