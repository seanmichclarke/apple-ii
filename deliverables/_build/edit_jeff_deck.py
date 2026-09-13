#!/usr/bin/env python3
"""Apply the proposed edits to Jeff Robison's deck.

The original (deck/Applesoft_Loaded_Dice_Deck_1.pptx) is never modified. Output goes to
deliverables/01_edited_deck/. Every change here is listed, with its reason and source,
in deliverables/01_edited_deck/EDIT_LOG.md.

Run from the project root:  .venv/bin/python deliverables/_build/edit_jeff_deck.py
"""
import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC = os.path.join(ROOT, 'deck', 'Applesoft_Loaded_Dice_Deck_1.pptx')
OUT = os.path.join(ROOT, 'deliverables', '01_edited_deck', 'Applesoft_Loaded_Dice_Deck_1_EDITED.pptx')

# Jeff's palette, read from the deck
PANEL = RGBColor(0x20, 0x22, 0x25)
GREEN = RGBColor(0x5F, 0xCB, 0x7E)
TEXT = RGBColor(0xE9, 0xE7, 0xE3)
MUTED = RGBColor(0x9D, 0xA3, 0x9B)
AMBER = RGBColor(0xE0, 0xA7, 0x3E)


# ---------------------------------------------------------------- helpers
def shape_by_id(slide, sid):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_id == sid:
                return sh
            if sh.shape_type == 6:
                found = walk(sh.shapes)
                if found is not None:
                    return found
        return None
    found = walk(slide.shapes)
    assert found is not None, 'shape id %d not found' % sid
    return found


def set_paras(shape, paras):
    """Replace all paragraphs, reusing the first paragraph's pPr and first run's rPr."""
    body = shape.text_frame._txBody
    ps = body.findall(qn('a:p'))
    ppr = ps[0].find(qn('a:pPr'))
    run = ps[0].find(qn('a:r'))
    rpr = run.find(qn('a:rPr')) if run is not None else None
    for p in ps:
        body.remove(p)
    for text in paras:
        p = etree.SubElement(body, qn('a:p'))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn('a:r'))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn('a:t')).text = text


def set_para_text(p_el, text):
    """Replace one paragraph's runs with a single run, keeping the first run's formatting."""
    runs = p_el.findall(qn('a:r'))
    rpr = runs[0].find(qn('a:rPr')) if runs else None
    rpr = copy.deepcopy(rpr) if rpr is not None else None
    for child in list(p_el):
        if child.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p_el.remove(child)
    r = etree.Element(qn('a:r'))
    if rpr is not None:
        r.append(rpr)
    etree.SubElement(r, qn('a:t')).text = text
    end = p_el.find(qn('a:endParaRPr'))
    if end is not None:
        end.addprevious(r)
    else:
        p_el.append(r)


def add_text(slide, x, y, w, h, paras, size, color, bold=False, font='Arial',
             align=None, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = t
        f = r.font
        f.name, f.size, f.bold = font, Pt(size), bold
        f.color.rgb = color
    return tb


def add_panel(slide, x, y, w, h, fill=PANEL, outline=None, dash=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if outline is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = outline
        shp.line.width = Pt(1.5)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    shp.shadow.inherit = False
    return shp


def next_id(slide):
    ids = [int(e.get('id')) for e in slide.shapes._spTree.iter(qn('p:cNvPr'))]
    return max(ids) + 1


def clone_shape(slide_from, sid, slide_to, dy=0.0):
    el = copy.deepcopy(shape_by_id(slide_from, sid)._element)
    el.find('.//' + qn('p:cNvPr')).set('id', str(next_id(slide_to)))
    slide_to.shapes._spTree.append(el)
    new = slide_to.shapes[-1]
    if dy:
        new.top = new.top + Inches(dy)
    return new


# ---------------------------------------------------------------- notes
NOTES = {
1: """Name and title. Fifteen seconds. No bio slide.

Before this slide: POWER-CYCLE the machine (power switch off, wait, on), then PRINT RND(1). Do it three times. Not Ctrl-Reset and not PR#6: a warm restart leaves the last seed byte in RAM and prints a different number (Aldridge 1987; Sander-Cederlof, AAL May 1984). Rehearse on the exact machine you bring.

Applesoft is Microsoft's 6502 BASIC, licensed by Apple. Keep that in mind; it matters on slide 5.""",

2: """This is the number they just watched three times. Don't explain it yet: "Hold onto that number."

Only a power-on repeats. Rerun the program, or reboot DOS, and the numbers look random. That's how this hid for years.

If someone says their machine prints a different number: the ROM copies only four of the five seed bytes. The fifth ($CD) is whatever RAM held. $FF or $FE gives .973136996; $00 gives .270011996. Each machine still repeats its own number every power-on. Sources: Sander-Cederlof, AAL May 1984; confirmed by running the ROM code.""",

3: """Thirty seconds. This room knows what a PRNG is.

The bottom box is the part that has to land: the seed decides where the sequence starts, and the formula decides what it looks like. This talk is about the seed. The Integer BASIC slide depends on it.""",

4: """Say out loud that $4E/$4F is not a jiffy clock. That's a Commodore term. The Apple II has no timer doing this.

This isn't a toy: a 15-bit maximal shift register, period 32,767, in 50 hand-assembled bytes. Add the 6 bytes in KEYIN and the whole RNG is 56 bytes.

The counter at $4E/$4F isn't just read as a seed. It IS the generator's state: RND scrambles it and writes it back, so every key wait stirs it.

Don't make the regression argument yet. Just show what worked.

If asked "was it actually good?": for games, yes. Not for statistics. RND(X) is the state MOD X, and every program shares the same 32,767-step cycle.""",

5: """Walk the left side first. Point at the bne going back to KEYIN. That's the whole entropy source.

The counter gets bumped before the keyboard is read, every pass. It's a loop count, not a keystroke count. The loop is 15 CPU cycles, so about 68,000 counts a second, and the 16-bit counter wraps roughly once a second.

Then the right side: Integer BASIC's RND reads $4E/$4F and writes it back.

Say the other three listings in one line ("same counter in the Autostart ROM, the IIe and the IIc"). Don't walk each one.

If asked: it only moves while a KEYIN-style routine waits for a key. A game that polls $C000 directly never moves it.

Sources: 6502disassembly.com/a2-rom/OrigF8ROM.html and IntegerBASIC.html. The monitor listing credits S. Wozniak and A. Baum. Woz wrote Integer BASIC with no assembler, so there is no original source, only hand-written pages in a binder. The Integer BASIC disassembly is Paul Santa-Maria's work, converted by Andy McFadden.""",

6: """This is the whole thing. Twenty-eight instructions. Microsoft wrote it; Apple shipped it unchanged in every Applesoft ROM through the enhanced IIe.

Point at the top: the seed comes from $C9, which got four of its five bytes from ROM at power-on. Point at the bottom: the answer goes back to $C9. Nothing else feeds it.

The comments on the listing ("very poor RND algorithm", "this does nothing") are Bob Sander-Cederlof's annotations, not Microsoft's. Say so before someone else does.

If asked about the add: the addend is so small it's lost to precision. Running the ROM code, it changed the stored seed in 5 of 57,021 steps. "Effectively nothing", not "nothing".

This listing is from the II Plus ROM (1979). The 1978 Applesoft releases loaded into RAM at other addresses.

Source: McFadden's SourceGen conversion of S-C DocuMentor, 6502disassembly.com/a2-rom/Applesoft.html. Microsoft's own source (github.com/microsoft/BASIC-M6502) has the same routine.""",

7: """Read the RND(n) line aloud: "generates a new random number each time it is used." True. It never says the sequence starts in the same place every power-on.

RND(-n): Apple documents repeatability as a feature, for debugging.

The manual tells you how to get the same sequence on purpose. It never tells you you're getting it by accident.

Hold this for the fix slide: when Aldridge's lab called Apple, Apple's advice was X = RND(-1*(PEEK(78)+256*PEEK(79))).

[Jeff: confirm which printing these scans come from before putting an edition or part number on screen.]""",

8: """Four decades of people finding the same thing and working around it one program at a time.

Lead with Aldridge: a memory experiment gave each day's first subject the same "random" word list, because that subject was tested right after the computer was turned on. That's the damage.

202 was measured on real Applesoft by two psychologists who needed randomized experiments. Sander-Cederlof saw repetition at the 37,758th number two years later. Same generator, different uncopied byte: both loops reproduce when you run the ROM code with different values in $CD.

Kaner sent you the paper directly. Say that out loud; it doesn't need screen space.

The Call-A.P.P.L.E. author is sometimes given as D. Sparks. That name isn't verified, so don't say it.""",

9: """Somebody will say just seed it first. Say it yourself before they do, and agree with it for new code. Apple told Aldridge's lab exactly that in 1987.

Then: you can't edit software you didn't write, the failure is silent so only people who already know can fix it, and this machine used to do it automatically.

HFIND: the S-C DocuMentor listing marks it "not called by any Applesoft routine", and there is no JSR or JMP to $F5CB anywhere in the II Plus ROM. HFIND runs $F5CB-$F5FF, 53 bytes. $F600 is an RTS that HLIN branches to (BEQ at $F59C), so the patch must be 53 bytes or less or HPLOT TO breaks. Know your patch's byte count. The only thing left that can break is a program that does CALL 62923 itself.

Expect: "What does RND(1) do twice in a tight loop with no keypress?" and "Does it work under ProDOS?" Have the answers ready.

Limits to state plainly if asked: it fixes the seed, not the generator; it needs a language card; it only gets keyboard timing if something waited for a key before the first RND.""",

10: """Takeaways first, then the link and QR. Leave this up during questions.

Questions to expect:
- Does the C64 have this too? Same Microsoft generator, same constants. But Microsoft's own source already made RND(0) read the PET's hardware timers and added TI, in the first PET ROM in 1977; Commodore carried that forward to the C64. On a Commodore the idiom is RND(-TI); on the Apple it's RND(-PEEK(78)-256*PEEK(79)).
- Why not replace the generator? Speed, no space, and RND(-n) has to keep repeating.
- What's at $F5CB? HFIND. Applesoft never calls it. 53 bytes. Answer it straight.
- My machine printed a different number. One seed byte is never set, so your machine repeats its own number.
- Does the patch fix the 202 loop? No. It fixes the seed.""",
}


def main():
    prs = Presentation(SRC)
    s = {i + 1: sl for i, sl in enumerate(prs.slides)}

    # 1. Badge numbers: 7 -> 6, 9 -> 8, add 9 to the closing slide.
    set_paras(shape_by_id(s[7], 3), ['6'])
    set_paras(shape_by_id(s[9], 3), ['8'])
    clone_shape(s[9], 2, s[10])
    badge = clone_shape(s[9], 3, s[10])
    set_paras(badge, ['9'])

    # 2. Slide 2: put the source for the claim on screen.
    add_text(s[2], 1.2, 4.75, 10.9, 0.9,
             ['“…exactly the same sequence each time the machine is powered on.”',
              'Aldridge, Behavior Research Methods, Instruments, & Computers, 1987'],
             15, MUTED, align=PP_ALIGN.CENTER)

    # 3. Slide 3: separate seed entropy from generator quality.
    set_paras(shape_by_id(s[3], 12),
              ['Entropy decides where the sequence starts. The formula decides what it looks like.'])

    # 4. Slide 4: measured claims instead of undefined ones.
    set_paras(shape_by_id(s[4], 7), [
        '15-bit shift register: repeats every 32,767 calls',
        'Integers only: RND(X) is the state MOD X',
        'Good enough for games, its stated purpose',
        '56 bytes in all: 6 in KEYIN, 50 in RND',
    ])
    set_paras(shape_by_id(s[4], 9), ['Seeded by you'])
    set_paras(shape_by_id(s[4], 10), [
        "There's a counter at $4E and $4F that goes up while the monitor sits waiting for a "
        "keypress. Integer BASIC's generator is that counter: RND scrambles it and writes it back."])

    # 5. Slide 5: name the ROMs actually shown instead of "all".
    claim = shape_by_id(s[5], 15)
    set_paras(claim, ['Same counter in the II, II Plus, IIe and IIc ROMs'])
    claim.width = Inches(5.9)

    # 6. Slide 6: name Microsoft, credit the comments, fix the plain-English description.
    set_paras(shape_by_id(s[6], 4), ["Applesoft II: Microsoft's RND"])
    sub = shape_by_id(s[6], 5)
    sub.text_frame.text = "Microsoft's 6502 BASIC, licensed by Apple"
    f = sub.text_frame.paragraphs[0].runs[0].font
    f.name, f.size, f.color.rgb = 'Arial', Pt(16), MUTED
    set_paras(shape_by_id(s[6], 8), ['Seed at $C9, from ROM'])
    set_paras(shape_by_id(s[6], 9), ['$4E/$4F still counts. Applesoft never reads it.'])
    set_paras(shape_by_id(s[6], 12), [
        'Takes the seed at $C9, multiplies, adds a constant too small to matter, swaps the top '
        'and bottom bytes, forces the result under 1, writes it back to $C9.'])
    add_text(s[6], 0.7, 6.47, 11.9, 0.45,
             ['II Plus ROM. Comments: Bob Sander-Cederlof (S-C DocuMentor), via Andy McFadden, '
              '6502disassembly.com. Code: Microsoft BASIC M6502.'],
             11, MUTED)
    pic = shape_by_id(s[6], 6)
    pic._element.nvPicPr.cNvPr.set(
        'descr', 'Applesoft II RND routine, $EFAE-$EFE7, from the S-C DocuMentor disassembly '
                 'via 6502disassembly.com')

    # 7. Slide 7: say what the scans prove.
    set_paras(shape_by_id(s[7], 4), ['The Applesoft Manual'])
    add_text(s[7], 9.44, 3.05, 3.4, 0.4, ['…starting here, every power-on.'], 13, GREEN)
    add_panel(s[7], 9.44, 3.9, 3.3, 1.75)
    add_text(s[7], 9.6, 4.0, 3.0, 1.55,
             ['Repeatable on purpose: documented.', 'Repeatable by accident: not.'],
             15, TEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # 8. Slide 8: correct citations, recover the 202 figure, add a sixth row.
    set_paras(shape_by_id(s[8], 4), ['Found, reported, worked around'])
    clone_shape(s[8], 13, s[8], dy=0.72)
    clone_shape(s[8], 14, s[8], dy=0.72)
    rows = [
        (5, 6, 'Kaner & Vokey, MICRO 1984',
         'Written 1982. RND fell into an endless loop of 202 numbers.'),
        (7, 8, 'Call-A.P.P.L.E., Jan 1983',
         '“RND is Fatally Flawed,” pp. 29–34.'),
        (9, 10, 'Sander-Cederlof, AAL May 1984',
         'Seed setup copies 4 of 5 bytes (fix: $F151 $1C→$1D). '
         'Repetition starts at the 37,758th number.'),
        (11, 12, 'Aldridge, BRMIC 1987',
         'A memory experiment gave each day’s first subject the same “random” word list.'),
        (13, 14, 'Gleason, 1988',
         'Collegiate Microcomputer. Tested Apple IIe RND; suggested seeds that pass.'),
    ]
    for left, right, a, b in rows:
        set_paras(shape_by_id(s[8], left), [a])
        set_paras(shape_by_id(s[8], right), [b])
    new_left, new_right = list(s[8].shapes)[-2:]
    set_paras(new_left, ['Empson, GS WorldView 1999'])
    set_paras(new_right, ["Write-up of Robert Moore's 1989 shift-register generator, seeded from $4E/$4F."])

    # 9. Slide 9: drop the hedge, state the scope.
    set_paras(shape_by_id(s[9], 4), ['The Fix: Put the Counter Back'])
    steps = shape_by_id(s[9], 19)
    for p in steps.text_frame._txBody.findall(qn('a:p')):
        txt = ''.join(t.text or '' for t in p.iter(qn('a:t')))
        if 'HFIND' in txt:
            set_para_text(p, 'Writes the patch code into the LC at $F5CB, over HFIND (53 bytes), '
                             'which no Applesoft routine calls.')
    set_paras(shape_by_id(s[9], 24), [
        'Keeps RND(−n) and RND(0) working as documented; positive arguments draw on the KEYIN counter.',
        'Fixes the seed, not the generator. Needs a language card.',
    ])

    # 10. Slide 10: takeaways, link, QR.
    add_text(s[10], 0.9, 0.75, 10.6, 1.55, [
        '1977: Integer BASIC seeded RND from you. Then Applesoft seeded it from ROM.',
        'It hid for years: rerun looked random; only power-on repeated.',
        'The fix puts the counter back without changing a single program.',
    ], 18, TEXT)
    add_text(s[10], 0.9, 3.8, 8.6, 0.5, ['[ URL: Jeff to supply ]'], 20, AMBER, font='Courier New')
    add_panel(s[10], 10.3, 3.6, 2.0, 2.0, fill=PANEL, outline=AMBER, dash=True)
    add_text(s[10], 10.3, 3.6, 2.0, 2.0, ['QR', 'Jeff to supply'], 13, AMBER,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Speaker notes for every slide.
    for n, text in NOTES.items():
        s[n].notes_slide.notes_text_frame.text = text

    cp = prs.core_properties
    cp.title = "Applesoft's Loaded Dice"
    cp.subject = 'The Truth About Randomness on the Apple II, Then and Now'
    cp.comments = 'Proposed edits for Jeff Robison to review. See EDIT_LOG.md.'

    prs.save(OUT)
    verify(OUT)


def verify(path):
    prs = Presentation(path)
    badges, problems = [], []
    for i, sl in enumerate(prs.slides, 1):
        ids = [e.get('id') for e in sl.shapes._spTree.iter(qn('p:cNvPr'))]
        if len(ids) != len(set(ids)):
            problems.append('slide %d has duplicate shape ids' % i)
        texts = []
        for e in sl.shapes._spTree.iter(qn('a:t')):
            texts.append(e.text or '')
        joined = ' '.join(texts)
        if 'supposedly unused' in joined:
            problems.append('slide %d still hedges on HFIND' % i)
        for sh in sl.shapes:
            if (sh.has_text_frame and sh.text_frame.text.strip()
                    and sh.left > Inches(11.8) and sh.top < Inches(0.6)):
                badges.append(sh.text_frame.text)
    notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides]
    if len(set(notes)) != len(notes):
        problems.append('duplicate speaker notes remain')
    print('saved', path)
    print('badges', badges)
    print('problems', problems or 'none')
    assert badges == [str(i) for i in range(10)], badges
    assert not problems


if __name__ == '__main__':
    main()
