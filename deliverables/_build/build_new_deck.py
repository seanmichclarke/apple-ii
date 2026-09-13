#!/usr/bin/env python3
"""Build the expanded deck from one slide definition, in four formats:

  02_new_deck/Applesoft_Loaded_Dice_EXPANDED.pptx   PowerPoint (opens in Keynote)
  02_new_deck/Applesoft_Loaded_Dice_EXPANDED.html   self-contained HTML deck (images inlined)
  02_new_deck/Applesoft_Loaded_Dice_EXPANDED_marp.md Marp / Markdown source
  02_new_deck/KEYNOTE_OUTLINE.md                    per-slide outline with notes and assets

Layouts are turned into drawing primitives (rect, text, image) once, and both the PPTX and
HTML renderers draw the same primitives, so the two stay visually in step.

Run from the project root:  .venv/bin/python deliverables/_build/build_new_deck.py
"""
import base64
import html
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUTDIR = os.path.abspath(os.path.join(HERE, '..', '02_new_deck'))
ASSETS = os.path.join(OUTDIR, 'assets')
BASENAME = 'Applesoft_Loaded_Dice_EXPANDED'

# Jeff's palette, so the two decks read as one talk
BG, PANEL, RULE = '17181A', '202225', '3A3E42'
GREEN, TEXT, MUTED = '5FCB7E', 'E9E7E3', '9DA39B'
AMBER, BARGRAY = 'E0A73E', '6B7176'
SANS, MONO = 'Arial', 'Courier New'
W, H = 13.333, 7.5


# ================================================================ slide definitions
def S(kind, title, notes, **kw):
    d = dict(kind=kind, title=title, notes=notes.strip(), section='main')
    d.update(kw)
    return d


def B(kind, title, notes, **kw):
    d = S(kind, title, notes, **kw)
    d['section'] = 'backup'
    return d


SLIDES = [
S('title', "Applesoft's Loaded Dice", """
Name and title. Fifteen seconds.

Before this slide: power-cycle the machine (switch off, wait, on) and PRINT RND(1). Three times. Not Ctrl-Reset, not PR#6: a warm restart keeps the fifth seed byte in RAM and prints a different number.

Applesoft is Microsoft's 6502 BASIC, licensed by Apple. That matters in a few slides.
""", subtitle='The Truth About Randomness on the Apple II, Then and Now',
  byline='Jeff Robison   ·   VCF Midwest 21 2026',
  tag='Expanded edition · draft for Jeff Robison’s review'),

S('hero', 'The Problem', """
The number they just watched three times. Don't explain it yet: "Hold onto that number."

If someone says their machine prints a different number: one seed byte is never set. $FF or $FE gives .973136996, $00 gives .270011996. Each machine still repeats its own number. We'll get to exactly why.
""", mono=[']PRINT RND(1)', '.973136996'],
  sub=['Power off. Power on. Same number.',
       '“…exactly the same sequence each time the machine is powered on.”'],
  source='Aldridge, Behavior Research Methods, Instruments, & Computers 19(4):397–399, 1987'),

S('cols', 'Why nobody noticed', """
This is why it survived for years. The programmer at the desk reruns, reboots DOS, hits Reset: all different. Only the first user of the day, after a cold power-on, gets the repeat.

Aldridge: "It is only when a machine is powered off and back on that the sequence begins repeating itself."

Ctrl-Reset leaves $C9-$CD untouched (measured by running the ROM's warm-start path).
""", lede='Everything a programmer does at a desk looks random.',
  cols=[('RUN it again', ['Different numbers.', 'The seed carries on from the last call.']),
        ('Reboot DOS, Ctrl-Reset', ['Different numbers.', 'The seed survives in RAM.']),
        ('Power switch', ['The same sequence.', 'The first user of the day sees it.'])],
  hl_col=2,
  source='Aldridge 1987; Ctrl-Reset behavior measured on the II Plus ROM warm-start path'),

S('cols', 'Pseudo-random means repeatable', """
Thirty seconds. This room knows what a PRNG is.

The split matters for the rest of the talk and for Q&A: a random seed doesn't make a generator's output patternless, and a good generator doesn't help if it always starts in the same place. This talk is about the seed. The generator is weak too, and we'll say so honestly.
""", cols=[('The seed', ['Where the sequence starts.', 'Needs entropy: something you can’t predict.']),
           ('The generator', ['What the sequence looks like.', 'Needs a long period and no visible pattern.'])],
  callout='Same seed → same numbers. Every time. This talk is about the seed.'),

S('cols', 'Integer BASIC, 1977: seeded by you', """
Say out loud that $4E/$4F is not a jiffy clock. The Apple II has no timer doing this.

Not a toy: a 15-bit maximal shift register in 50 hand-assembled bytes. Woz wrote Integer BASIC with no assembler; there is no source file, only hand-written pages.

The counter isn't just read as a seed. It IS the generator's state: RND scrambles $4E/$4F and writes it back, so every key wait stirs it.

Don't make the regression argument yet. Just show what worked.
""", lede='Woz’s generator at $EF4E, hand-assembled.',
  cols=[('The generator', ['15-bit shift register', 'Repeats every 32,767 calls', 'RND(X) = state MOD X', '50 bytes']),
        ('The seed', ['$4E/$4F counts while the monitor waits for a key',
                      'That counter is the generator’s state', 'Every key wait stirs it', '6 bytes, in KEYIN'])],
  callout='Good enough for games, its stated purpose. Not for statistics.',
  source='Integer BASIC disassembly (Santa-Maria, via McFadden); Sander-Cederlof, AAL Aug 1981; Wozniak, BYTE May 1977; period measured by emulation'),

S('image', 'The entropy source is a busy-wait', """
Point at the bne back to KEYIN. That's the whole entropy source.

The counter is bumped before the keyboard is read, every pass. A loop count, not a keystroke count. 15 CPU cycles per pass: about 68,000 counts a second, wrapping roughly once a second. Human reaction time jitters by far more than a second's worth of counts, which is why it works.

Same counter in the Autostart ROM, the IIe and the IIc. Say that in one line.

Caveat if asked: it only moves while a KEYIN-style routine waits. A game polling $C000 never moves it.
""", lede='Not a timer. Not a jiffy clock. A loop count.',
  image='keyin_orig_f8_fd1b.png', caption='Monitor KEYIN, $FD1B. Listing credited to S. Wozniak and A. Baum.',
  side=[('Rate', ['15 CPU cycles per pass', '≈ 68,000 counts a second', 'Wraps ≈ once a second']),
        ('Same counter in', ['Autostart F8 ($FD1B)', 'IIe 80-col ($CB15)', 'IIc ($CC71)']),
        ('Only moves', ['while a KEYIN-style routine waits for a key'])],
  source='6502disassembly.com: OrigF8ROM, AutoF8ROM, Unenh_IIe_80col, IIc_16kb; Apple II Reference Manual p. 32'),

S('manual', 'What the manual promised', """
The promise before the reality.

Read the RND(n) line aloud: "generates a new random number each time it is used." True. It never says the sequence starts in the same place every power-on.

RND(-n): Apple documents repeatability as a feature, for debugging.

The manual tells you how to get the same sequence on purpose. It never tells you you're getting it by accident.

[Jeff: confirm which printing these scans come from.]
""", rows=[('RND(n)', 'manual_rnd_positive.png'), ('RND(−n)', 'manual_rnd_negative.png'),
           ('RND(0)', 'manual_rnd_zero.png')],
  shot='screen_973136996.png', shot_caption='…starting here, every power-on.',
  callout=['Repeatable on purpose: documented.', 'Repeatable by accident: not.'],
  source='Applesoft II BASIC Programming Reference Manual, Apple Computer (printing to confirm)'),

S('image', "Applesoft II: Microsoft's RND", """
The reveal. Twenty-eight instructions. Microsoft wrote it; Apple shipped it unchanged in every Applesoft ROM through the enhanced IIe.

Top: the seed comes from $C9. Bottom: the answer goes back to $C9. Nothing else feeds it.

The comments on this listing are Bob Sander-Cederlof's annotations, not Microsoft's. Say so before someone else does.

The add is lost to precision: running the ROM, it changed the stored seed in 5 of 57,021 steps. "Effectively nothing."
""", lede='Microsoft wrote it. Apple shipped it unchanged through the enhanced IIe.',
  image='applesoft_rnd_efae.png',
  caption='II Plus ROM. Comments: Bob Sander-Cederlof (S-C DocuMentor), via Andy McFadden. Not Microsoft’s.',
  side=[('Seed at $C9', ['Filled from ROM at power-on.', 'Nothing outside RND touches it.']),
        ('$4E/$4F', ['Still counts.', 'Applesoft never reads it.'], AMBER)],
  source='Microsoft BASIC M6502 v1.1 source, github.com/microsoft/BASIC-M6502; 6502disassembly.com/a2-rom/Applesoft.html'),

S('split', 'One byte short', """
The most interesting detail in the story, and it explains why your number might differ from mine.

The copy loop moves CHRGET and the seed into zero page. The count is one too small. Four seed bytes arrive; the fifth, $CD, is whatever RAM held.

The bug is in Microsoft's source. But look at $F157: STX SPEEDZ is Apple's own instruction, for SPEED=, inserted inside this loop. Apple's engineers worked inside this loop and kept the short count.

Sander-Cederlof published the one-byte fix in 1984: $F151 from $1C to $1D.

$CD at power-on on real DRAM is not published anywhere. [Jeff: measure on the demo machine; backup B3.]
""", lede='Cold start copies 4 of the 5 seed bytes. The fifth, $CD, is whatever RAM held.',
  code=['F123: 80 4F C7 52 58   seed table',
        '',
        'F150: A2 1C      LDX #$1C     ; one short',
        'F152: BD 0A F1   LDA $F10A,X',
        'F155: 95 B0      STA $B0,X',
        'F157: 86 F1      STX SPEEDZ   ; Apple’s',
        'F159: CA         DEX',
        'F15A: D0 F6      BNE $F152'],
  header=['$CD at power-on', 'First RND(1)'],
  rows=[['$FE or $FF', '.973136996'], ['$00', '.270011996'],
        ['$58 (intended)', '.512199496'], ['$AA', '.738761996']],
  hl=0,
  callout='256 possible bytes → 181 different first numbers. Each machine still repeats its own.',
  source='II Plus ROM bytes; S-C DocuMentor; Sander-Cederlof, AAL May 1984; values from running the ROM code'),

S('bars', 'Two published loops, one generator', """
Kaner and Vokey saw a 202-number loop. Sander-Cederlof saw repetition at 37,758. Both were right.

Running every possible value of the uncopied byte: nearly a quarter of cold starts end up in the 202-number loop. Microsoft's intended byte, $58, is one of them.

Reseeding doesn't escape it: RND(-52894) and RND(-22258) both land in the 202 loop. Keep that for the question "does the patch fix the loop?" (No.)

The four 12,559 results weren't checked to be one loop. Don't dwell on that bar.
""", lede='Kaner & Vokey saw a loop of 202. Sander-Cederlof saw 37,758. Both were right.',
  chart_title='Where all 256 possible cold starts end up, by loop length',
  bars=[('37,758', 110), ('32,366', 78), ('202', 59), ('4,082', 5), ('12,559', 4)],
  hl_bar='202',
  side=[('$CD = $FF', ['37,758-number loop', 'after 6,818 calls']),
        ('$CD = $58, the intended byte', ['202-number loop', 'after 15,382 calls'], AMBER),
        ('Reseeding', ['RND(−52894) lands in the 202 loop too'])],
  source='Kaner & Vokey, MICRO June 1984; Sander-Cederlof, AAL May 1984; all 256 cold starts run on the II Plus ROM'),

S('table', 'A regression of the default', """
The strongest honest version of the thesis. Concede that both generators are weak.

Integer BASIC: one 32,767-long cycle; a seed only picks a position. State MOD X. Weak.
Applesoft: bigger state, short loops. Also weak.

The difference that matters is the default. Integer BASIC got keypress timing automatically, because its state and the counter are the same bytes. Applesoft kept the counter running and ignored it.
""", lede='Both generators are weak. Only one seeded itself.',
  header=['', 'Integer BASIC (1977)', 'Applesoft II (Microsoft)'],
  rows=[['State', '15 bits at $4E/$4F', '5-byte float at $C9–$CD'],
        ['Loop length', '32,767, one cycle', '202 to 37,758, depends on seed'],
        ['Seed by default', 'Keypress timing, automatically', 'ROM, plus one leftover byte'],
        ['Reads $4E/$4F', 'It is $4E/$4F', 'Never'],
        ['Programmer must', 'Nothing', 'RND(−(PEEK(78)+256*PEEK(79)))']],
  hl=2, colw=[0.2, 0.37, 0.43],
  callout='The machine didn’t lose its entropy source in 1978. The default stopped using it.'),

S('bullets', 'Why 1978 got worse', """
No record explains the decision. Every statement about motive is inference.

What is documented: Microsoft's source builds RND two ways. REALIO=3, the Commodore PET, reads free-running VIA timers in RND(0). Everyone else, including Apple, gets the fixed-seed version.

Apple wasn't locked out. Applesoft II's cold start has Apple's own insertions: SPEED=, TRACE, the silent RAM probe. STX SPEEDZ sits inside the seed-copy loop.

Say "nobody connected it", not "nobody knew" or "Microsoft was careless".
""", bullets=['Microsoft’s BASIC had two RNDs: timer-seeded for Commodore, a fixed seed for everyone else.',
              'Apple got the everyone-else build.',
              'Apple edited code inside the seed-copy loop (STX SPEEDZ at $F157) and kept the short count.',
              'The same RND bytes shipped in the II Plus, IIe and enhanced IIe ROMs.'],
  callout='No record explains why. Nobody connected the counter the monitor was already running.',
  source='Microsoft m6502.asm (IFE REALIO-3); Steil, msbasic; ROM comparison of II Plus, IIe, enhanced IIe; Weyhrich, Apple II History ch. 16'),

S('quote', 'The damage', """
This is the stakes. Thirty seconds, and read it slowly.

A psychology lab ran a memory experiment in which every subject was supposed to get an individually randomized word list. The same ordering kept reappearing, and it was always the first subject of the day, right after the computer was switched on.

Kaner and Vokey were writing because their experiments needed randomization too.

Don't name a game that shipped broken unless you have evidence for it. We don't.
""", quote='…the subjects receiving the repeated lists were those tested at the beginning of each day, '
           'immediately after the computer had been turned on.',
  attribution='Aldridge, 1987: a memory experiment with “individually randomized” word lists',
  lines=['Kaner & Vokey needed RND to randomize experiments too.', 'Published research ran on this generator.'],
  source='Aldridge, Behavior Research Methods, Instruments, & Computers 19(4):397–399, 1987'),

S('table', 'Found, reported, worked around', """
Four decades of people finding the same thing and working around it one program at a time.

Kaner sent you the paper directly; say that aloud.

Call-A.P.P.L.E.'s author is sometimes given as D. Sparks. That's unverified, so it isn't on the slide.

Gleason: only the abstract has been read. Don't quote numbers from it.
""", header=['Year', 'Who', 'What'],
  rows=[['1982/84', 'Kaner & Vokey, MICRO', 'Loop of 202 numbers; replacement generator'],
        ['1983', 'Call-A.P.P.L.E.', '“RND is Fatally Flawed,” pp. 29–34'],
        ['1984', 'Sander-Cederlof, AAL', 'Seed copies 4 of 5 bytes; repetition at 37,758'],
        ['1987', 'Aldridge, BRMIC', 'Same word lists every morning; Apple’s advice: seed from $4E/$4F'],
        ['1988', 'Gleason, Collegiate Micro.', 'Statistical tests; suggested seeds'],
        ['1989/99', 'Moore; Empson', 'Shift-register replacement seeded from $4E/$4F']],
  colw=[0.13, 0.3, 0.57], rowh=0.52,
  callout='Found in 1982. Worked around, one program at a time. Never fixed in ROM.'),

S('cols', '“Just seed it first”', """
Say it yourself before they do, and agree with it for new code.

Then the three-part answer. You can't edit software you didn't write. The failure is silent, so only people who already know can fix it. And this machine used to do it automatically.

Edge case if asked: if both counter bytes are zero, RND(0) doesn't reseed. One in 65,536.
""", lede='Right. That’s what Apple told Aldridge’s lab in 1987.',
  cols=[('New code', ['`X = RND(−(PEEK(78)+256*PEEK(79)))`', 'after a key wait, before the first RND']),
        ('Existing code', ['You can’t edit software you didn’t write.',
                           'The failure is silent: only people who already know can fix it.',
                           'This machine used to do it automatically.'])],
  callout='The patch is for the programs nobody will ever edit.',
  source='Aldridge 1987 (Apple’s advice to the lab)'),

S('memmap', 'The fix: put the counter back', """
Copy the whole ROM into the language card, patch the copy, and read from the card from then on. RND's entry jumps to patch code that lives where HFIND was.

HFIND: 53 bytes, $F5CB to $F5FF. The S-C DocuMentor marks it "not called by any Applesoft routine", and no JSR or JMP to $F5CB exists anywhere in the ROM.

$F600 is an RTS that HLIN branches to. The patch must stop at $F5FF, or HPLOT TO breaks.

[Jeff: state the patch's byte count here.]
""", lede='Copy the ROM into the language card, patch the copy, read from the card.',
  regions=[('$F800–$FFFF', 'Monitor, copied', 'copy'),
           ('$F601', 'DRAW, untouched', 'keep'),
           ('$F600', 'RTS shared with HLIN, untouched', 'keep'),
           ('$F5CB–$F5FF', 'Patch code over HFIND: 53 bytes, never called by Applesoft', 'patch'),
           ('$EFAE', 'RND entry: JMP $F5CB', 'patch'),
           ('$D000–$F7FF', 'Applesoft, copied', 'copy')],
  side=[('Loader steps', ['1. Write-enable the LC, ROM still readable', '2. Copy $D000–$FFFF into the LC',
                          '3. Write the patch at $F5CB', '4. Write JMP $F5CB at $EFAE', '5. Switch the LC to read'])],
  source='Jeff Robison, LC_Loader / Patch_lc; HFIND per S-C DocuMentor; byte search of the II Plus ROM'),

S('cols', 'What the fix does, and doesn’t', """
Scope stated before anyone asks.

It fixes the seed, not the generator. The 202 loop is still reachable. Replacing the generator would break the documented RND(-n) repeatability and there's no room.

It needs a key wait before the first RND to get any timing. Aldridge flagged the same limit for Apple's own advice.

[Jeff: fill in the exact behavior of positive RND, and the compatibility results from backup B4, before this slide goes on stage.]
""", cols=[('Does', ['Positive RND draws on the KEYIN counter', 'RND(−n) still repeats; RND(0) still replays',
                     'No program changes']),
           ('Doesn’t', ['Fix the generator: the 202 loop is still there', 'Help before the first key wait',
                            'Work without a language card'])],
  callout='[Jeff: confirm behavior and compatibility (backup B4) before this slide is shown]', callout_amber=True),

S('hero', 'Proof', """
Symmetry with the opening: the talk opened with proof of the problem, so close with proof of the fix. About 90 seconds.

Power on, load the patch, PRINT RND(1). Power-cycle, load it again, PRINT RND(1). Different numbers.

Typing the BRUN command is itself a key wait, so the counter has moved.

If power-cycling on stage is risky, use a recorded clip. [Jeff: match the loader's file name.]
""", mono=['[power off, power on]', ']BRUN LC_LOADER', ']PRINT RND(1)', '', '[power off, power on]',
           ']BRUN LC_LOADER', ']PRINT RND(1)'],
  mono_size=22,
  sub=['Two power-ons. Two different numbers.', 'Live, or a recorded clip if power-cycling on stage is risky.']),

S('cols', 'Now: the same shape, 2006', """
Present as a parallel, not a lineage. Nobody learned or failed to learn this from the Apple II.

Debian, September 2006: a change to OpenSSL's random number generator, made to silence a memory-checker warning, removed nearly all entropy mixing. Keys generated on Debian and Ubuntu were predictable until the May 2008 disclosure. CVE-2008-0166.

Same shape as 1978: the entropy input was silently disconnected, and the output still looked random.

Lighter version: C's rand() without srand() behaves as if seeded with 1. Same sequence, every run, today.
""", lede='A change silently disconnects the entropy. The output still looks random.',
  cols=[('Debian OpenSSL, 2006–2008', ['A Debian change removed nearly all entropy from OpenSSL’s PRNG',
                                          'Keys were predictable until May 2008', 'CVE-2008-0166']),
        ('C, today', ['rand() without srand() starts as if seeded with 1', 'Same sequence, every run'])],
  callout='A parallel, not a lineage.',
  source='Debian Security Advisory DSA-1571-1 (May 2008); ISO C, rand/srand'),

S('closing', 'Try it yourself', """
Takeaways first, then the link. Leave this up during questions.

Questions to expect are in QA_PREP.md: the different-number question, the C64, why not replace the generator, $F5CB, ProDOS, Ctrl-Reset, "does it fix the 202 loop?".
""", takeaways=['1977: Integer BASIC seeded RND from you. Applesoft seeded it from ROM.',
                'It hid for years: rerun looked random; only power-on repeated.',
                'The fix puts the counter back without changing a single program.'],
  byline='Jeff Robison   ·   VCF Midwest 21 2026'),

# ---------------------------------------------------------------- backup
B('code', 'Backup: Applesoft RND, byte by byte', """
For the hex-checkers. Every byte matches the AppleWin II Plus ROM, and the same bytes sit at the same addresses in the IIe and enhanced IIe ROMs.
""", code=['EFAE: 20 82 EB  JSR SIGN        ; -1 / 0 / +1',
           'EFB1: AA        TAX',
           'EFB2: 30 18     BMI $EFCC       ; negative: reseed from arg',
           'EFB4: A9 C9     LDA #$C9',
           'EFB6: A0 00     LDY #$00',
           'EFB8: 20 F9 EA  JSR LOAD_FAC    ; FAC = seed',
           'EFBB: 8A        TXA',
           'EFBC: F0 E7     BEQ $EFA5       ; RND(0): return seed',
           'EFBE: A9 A6     LDA #$A6',
           'EFC0: A0 EF     LDY #$EF',
           'EFC2: 20 7F E9  JSR FMULT       ; x 11879546.40625',
           'EFC5: A9 AA     LDA #$AA',
           'EFC7: A0 EF     LDY #$EF',
           'EFC9: 20 BE E7  JSR FADD        ; + 3.93E-8, lost',
           'EFCC: A6 A1     LDX FAC+4       ; swap highest and',
           'EFCE: A5 9E     LDA FAC+1       ;   lowest bytes',
           'EFD0: 85 A1     STA FAC+4',
           'EFD2: 86 9E     STX FAC+1',
           'EFD4: A9 00     LDA #$00',
           'EFD6: 85 A2     STA FAC_SIGN    ; force positive',
           'EFD8: A5 9D     LDA FAC',
           'EFDA: 85 AC     STA FAC_EXT     ; old exponent -> guard',
           'EFDC: A9 80     LDA #$80',
           'EFDE: 85 9D     STA FAC         ; value < 1',
           'EFE0: 20 2E E8  JSR NORMALIZE',
           'EFE3: A2 C9     LDX #$C9',
           'EFE5: A0 00     LDY #$00',
           'EFE7: 4C 2B EB  JMP STORE_FAC   ; round, write $C9-$CD'],
  side=[('Not a textbook LCG', ['No modulus.', 'Float multiply, byte swap, renormalize.',
                                'LCG lattice theory doesn’t transfer.']),
        ('Measured', ['Fraction bit 25 is set in 99.75% of outputs.',
                      'Leading digits pass uniformity and serial-pair χ².'])],
  source='AppleWin Apple2_Plus.rom (SHA-1 33a24f54…); S-C DocuMentor labels, simplified'),

B('table', 'Backup: the constants, as actually read', """
Each constant is stored as 4 bytes; the float loader always reads 5, so each picks up the next byte in ROM. The truncation is in Microsoft's own source (the RND constants were never widened for 5-byte floats).
""", header=['Constant', 'Stored at', 'Bytes actually read', 'Value used'],
  rows=[['Multiplier', '$EFA6', '98 35 44 7A 68', '11,879,546.40625'],
        ['Addend', '$EFAA', '68 28 B1 46 20', '3.93 × 10⁻⁸'],
        ['Seed', '$F123', '80 4F C7 52 (58 never copied)', '≈ .811635157 + leftover byte']],
  colw=[0.18, 0.14, 0.36, 0.32], mono_cols=[2],
  callout='The addend changed the stored seed in 5 of 57,021 steps. “Effectively nothing,” not “nothing.”',
  source='ROM bytes; Microsoft m6502.asm (octal); decoded with exact rational arithmetic; emulation of the II Plus ROM'),

B('table', 'Backup: $CD on real hardware (to fill in)', """
[Jeff: fill in before the show.] Power on, and before any RND: FOR I = 201 TO 205 : PRINT PEEK(I);" "; : NEXT, then PRINT RND(1).

Emulator reference: $FF or $FE gives .973136996, $00 gives .270011996, $58 gives .512199496.
""", header=['Machine / ROM', 'Boot path', 'PEEK(205)', 'First RND(1)'],
  rows=[['II Plus', 'No disk, Ctrl-Reset to ]', '', ''], ['II Plus', 'DOS 3.3', '', ''],
        ['IIe (enhanced)', 'DOS 3.3', '', ''], ['IIe (enhanced)', 'ProDOS / BASIC.SYSTEM', '', ''],
        ['IIc', 'ProDOS / BASIC.SYSTEM', '', '']],
  colw=[0.25, 0.35, 0.18, 0.22],
  callout='Emulator reference:  $FF → .973136996   ·   $00 → .270011996   ·   $58 → .512199496'),

B('table', 'Backup: patch compatibility (to fill in)', """
[Jeff: mark each cell tested on hardware, tested in an emulator (name it), or untested.] If the loader refuses a configuration, say so on slide 15.
""", header=['Configuration', 'Risk', 'Result'],
  rows=[['ProDOS 8 / BASIC.SYSTEM', 'ProDOS lives in LC RAM; loader overwrites it', 'untested'],
        ['DOS 3.3, FP / INT', 'LC switches flip; patch may silently vanish', 'untested'],
        ['Original II, Integer in ROM', 'Copies Integer BASIC, JMPs into its middle', 'untested'],
        ['IIe / IIc, Ctrl-Reset', 'RESET forces ROM read; patch off', 'untested'],
        ['Programs using the LC', 'Pascal, CP/M, RAM disks, /RAM', 'untested'],
        ['//c ROMs, IIc Plus, IIgs', 'Applesoft bytes not checked', 'untested']],
  colw=[0.3, 0.52, 0.18], rowh=0.55),

B('bullets', 'Backup: HFIND and $F5CB', """
The answer to "what's at $F5CB?". Answer it straight.
""", bullets=['HFIND runs $F5CB–$F5FF: 53 bytes. “Not called by any Applesoft routine” (S-C DocuMentor).',
              'No JSR or JMP to $F5CB anywhere in the II Plus ROM (byte search).',
              '$F600 is an RTS that HLIN reaches from BEQ at $F59C. Overwrite it and HPLOT TO breaks.',
              'DRAW starts at $F601.',
              'Residual risk: a program that does CALL 62923 itself.',
              'Regression test after patching: HPLOT … TO, DRAW, XDRAW, SCALE, ROT.'],
  source='S-C DocuMentor; byte search and branch-target search of AppleWin Apple2_Plus.rom'),

B('cols', 'Backup: Commodore, same generator, different integration', """
The correction to "Commodore wired RND(0) to timers": the timer code is in Microsoft's own source for the Commodore target, and it shipped in the first PET ROM in 1977. Commodore carried it to the C64, where it goes through the KERNAL.
""", cols=[('Same', ['Microsoft’s generator and constants', 'Seed-copy bug in every Microsoft 6502 BASIC']),
           ('Different', ['Microsoft’s Commodore build: RND(0) reads hardware timers (PET, 1977)',
                          'A TI clock variable', 'Idiom: `RND(−TI)`'])],
  callout='Commodore got the platform integration from Microsoft. Apple did its own, and skipped RND.',
  source='Microsoft m6502.asm (IFE REALIO-3, TIME==1); Steil, msbasic and c64ref'),

B('bullets', 'Backup: sources', """
Full list with read status in 03_collateral/SOURCES.md.
""", bullets=['Microsoft, BASIC M6502 8K v1.1 source, github.com/microsoft/BASIC-M6502 (2025)',
              'Sander-Cederlof, S-C DocuMentor: Applesoft; McFadden, 6502disassembly.com',
              'Kaner & Vokey, “A Better Random Number Generator for Apple’s Floating Point BASIC,” MICRO, June 1984',
              'Call-A.P.P.L.E., “RND is Fatally Flawed,” Jan 1983, pp. 29–34',
              'Sander-Cederlof, “Random Numbers for Applesoft,” Apple Assembly Line, May 1984',
              'Aldridge, “Cautions regarding random number generation on the Apple II,” BRMIC 19(4), 1987',
              'Gleason, Collegiate Microcomputer 6(2), 1988 (abstract read)',
              'Empson, GS WorldView, Nov 1999 (on Moore, 1989)',
              'Apple II Reference Manual (1978); Applesoft II BASIC Programming Reference Manual (1978)',
              'Steil, pagetable.com and msbasic; Weyhrich, Apple II History'],
  body_size=15),
]


# ================================================================ primitives
def P(t, size=16, color=TEXT, bold=False, font=SANS, after=6):
    return dict(t=t, size=size, color=color, bold=bold, font=font, after=after)


def rect(x, y, w, h, fill=PANEL, outline=None, dash=False, tip=None):
    return ('rect', x, y, w, h, fill, outline, dash, tip)


def text(x, y, w, h, paras, align='l', anchor='t'):
    return ('text', x, y, w, h, paras, align, anchor)


def img(x, y, w, h, asset, align='c'):
    return ('img', x, y, w, h, asset, align)


def line_para(t, size, color=TEXT, after=8):
    """Backtick-wrapped text renders as code."""
    if t.startswith('`') and t.endswith('`'):
        return P(t[1:-1], size - 1, GREEN, font=MONO, after=after)
    return P(t, size, color, after=after)


_size_cache = {}


def asset_size(name):
    if name not in _size_cache:
        _size_cache[name] = Image.open(os.path.join(ASSETS, name)).size
    return _size_cache[name]


def fit(name, w, h):
    pw, ph = asset_size(name)
    scale = min(w / pw, h / ph)
    return pw * scale, ph * scale


# ================================================================ layouts
def chrome(sp, badge):
    out = [rect(11.95, 0.42, 0.62, 0.62),
           text(11.95, 0.42, 0.62, 0.62, [P(badge, 15 if len(badge) < 3 else 12, GREEN, True, MONO, 0)], 'c', 'm')]
    tsize = 32 if len(sp['title']) <= 38 else 28
    out.append(text(0.7, 0.45, 11.0, 0.75, [P(sp['title'], tsize, TEXT, True, after=0)], anchor='m'))
    if sp.get('lede'):
        out.append(text(0.7, 1.25, 11.2, 0.5, [P(sp['lede'], 18, MUTED, after=0)]))
    if sp.get('source'):
        out.append(text(0.7, 6.9, 11.9, 0.4, [P(sp['source'], 10.5, MUTED, after=0)]))
    return out


def callout_box(x, y, w, h, lines, amber=False, size=19):
    col = AMBER if amber else TEXT
    return [rect(x, y, w, h), rect(x, y, 0.07, h, fill=AMBER),
            text(x + 0.35, y, w - 0.6, h, [P(t, size, col, bold=amber, after=4) for t in lines], anchor='m')]


def L_title(sp, badge):
    return [text(0.9, 2.3, 11.5, 1.0, [P(sp['title'], 52, TEXT, True, after=0)]),
            text(0.9, 3.35, 11.5, 0.5, [P(sp['subtitle'], 22, GREEN, after=0)]),
            rect(0.92, 4.2, 1.6, 0.03, fill=RULE),
            text(0.9, 4.5, 11.5, 0.4, [P(sp['byline'], 16, MUTED, after=0)]),
            text(0.9, 6.75, 11.5, 0.4, [P(sp['tag'], 12, MUTED, after=0)]),
            rect(11.95, 0.42, 0.62, 0.62), text(11.95, 0.42, 0.62, 0.62, [P(badge, 15, GREEN, True, MONO, 0)], 'c', 'm')]


def L_hero(sp, badge):
    out = chrome(sp, badge)
    size = sp.get('mono_size', 30)
    lh = size / 72 * 1.3
    ph = len(sp['mono']) * lh + 0.6
    pw = 7.2 if size < 30 else 5.2
    y = 1.75 if len(sp['mono']) > 3 else 2.2
    out.append(rect((W - pw) / 2, y, pw, ph))
    out.append(text((W - pw) / 2 + 0.4, y + 0.3, pw - 0.6, ph - 0.6,
                    [P(t, size, GREEN, font=MONO, after=0) for t in sp['mono']]))
    sub = [P(sp['sub'][0], 22, TEXT, True, after=10)] + [P(t, 18, MUTED, after=6) for t in sp['sub'][1:]]
    out.append(text(0.9, y + ph + 0.35, 11.5, 1.4, sub, 'c'))
    return out


def L_cols(sp, badge):
    out = chrome(sp, badge)
    cols = sp['cols']
    n = len(cols)
    gap = 0.3
    cw = (11.9 - gap * (n - 1)) / n
    top = 1.95
    ph = 3.35 if sp.get('callout') else 4.6
    for i, col in enumerate(cols):
        head, lines = col[0], col[1]
        x = 0.7 + i * (cw + gap)
        accent = sp.get('hl_col') == i
        out.append(rect(x, top, cw, ph, outline=AMBER if accent else None))
        paras = [P(head, 19, AMBER if accent else GREEN, True, after=12)]
        paras += [line_para(t, 17 if n < 3 else 16) for t in lines]
        out.append(text(x + 0.3, top + 0.25, cw - 0.6, ph - 0.4, paras))
    if sp.get('callout'):
        out += callout_box(0.7, top + ph + 0.2, 11.9, 1.05, [sp['callout']], sp.get('callout_amber', False))
    return out


def L_bullets(sp, badge):
    out = chrome(sp, badge)
    size = sp.get('body_size', 20)
    top = 1.95 if not sp.get('lede') else 1.95
    h = 3.5 if sp.get('callout') else 4.7
    paras = [P('▸  ' + t, size, TEXT, after=14 if size >= 18 else 7) for t in sp['bullets']]
    out.append(text(0.9, top, 11.5, h, paras))
    if sp.get('callout'):
        out += callout_box(0.7, 5.6, 11.9, 1.05, [sp['callout']])
    return out


def table_prims(x, y, w, header, rows, colw=None, rowh=0.5, hl=None, size=15, mono_cols=()):
    colw = colw or [1 / len(header)] * len(header)
    out = []
    cx = [x]
    for f in colw[:-1]:
        cx.append(cx[-1] + f * w)
    for j, hcell in enumerate(header):
        out.append(text(cx[j] + 0.15, y, colw[j] * w - 0.2, 0.42, [P(hcell, 13, MUTED, True, after=0)], anchor='m'))
    out.append(rect(x, y + 0.45, w, 0.015, fill=RULE))
    yy = y + 0.55
    for i, row in enumerate(rows):
        is_hl = hl == i
        out.append(rect(x, yy, w, rowh - 0.06))
        if is_hl:
            out.append(rect(x, yy, 0.07, rowh - 0.06, fill=AMBER))
        for j, cell in enumerate(row):
            first = j == 0
            color = AMBER if (is_hl and not first) else (GREEN if first else TEXT)
            font = MONO if j in mono_cols else SANS
            out.append(text(cx[j] + 0.18, yy, colw[j] * w - 0.25, rowh - 0.06,
                            [P(cell, size - (1 if font == MONO else 0), color, first or is_hl, font, 0)], anchor='m'))
        yy += rowh
    return out, yy


def L_table(sp, badge):
    out = chrome(sp, badge)
    top = 1.95
    prims, end = table_prims(0.7, top, 11.9, sp['header'], sp['rows'], sp.get('colw'),
                             sp.get('rowh', 0.58), sp.get('hl'), 16, sp.get('mono_cols', ()))
    out += prims
    if sp.get('callout'):
        out += callout_box(0.7, max(end + 0.2, 5.55), 11.9, 1.0, [sp['callout']])
    return out


def L_split(sp, badge):
    out = chrome(sp, badge)
    top = 1.95
    out.append(rect(0.7, top, 6.4, 3.35))
    out.append(text(0.95, top + 0.25, 6.0, 3.0,
                    [P(t if t else ' ', 15, GREEN if t.startswith('F150') else TEXT, font=MONO, after=2)
                     for t in sp['code']]))
    prims, _ = table_prims(7.35, top, 5.25, sp['header'], sp['rows'], [0.52, 0.48], 0.66, sp.get('hl'), 16, [1])
    out += prims
    out += callout_box(0.7, 5.5, 11.9, 1.05, [sp['callout']])
    return out


def side_panels(x, y, w, side, maxh):
    out = []
    for item in side:
        head, lines = item[0], item[1]
        accent = item[2] if len(item) > 2 else GREEN
        h = 0.62 + 0.3 * len(lines)
        out.append(rect(x, y, w, h))
        paras = [P(head, 15, accent, True, after=6)] + [P(t, 14, TEXT, after=3) for t in lines]
        out.append(text(x + 0.25, y + 0.18, w - 0.4, h - 0.25, paras))
        y += h + 0.18
    return out


def L_image(sp, badge):
    out = chrome(sp, badge)
    top = 1.95
    iw, ih = fit(sp['image'], 7.9, 4.3)
    out.append(img(0.7, top, 7.9, ih, sp['image'], 'l'))
    out.append(text(0.7, top + ih + 0.12, 7.9, 0.5, [P(sp['caption'], 11, MUTED, after=0)]))
    out += side_panels(8.85, top, 3.75, sp['side'], 4.8)
    return out


def L_manual(sp, badge):
    out = chrome(sp, badge)
    y = 1.5
    for label, asset in sp['rows']:
        iw, ih = fit(asset, 6.9, 3.0)
        out.append(text(0.7, y, 1.4, 0.4, [P(label, 16, GREEN, True, MONO, 0)]))
        out.append(rect(2.05, y - 0.05, 7.0, ih + 0.1, fill='F4F1EA'))
        out.append(img(2.1, y, 6.9, ih, asset, 'l'))
        y += ih + 0.4
    sw, sh = fit(sp['shot'], 3.3, 1.0)
    out.append(img(9.35, 1.55, 3.3, sh, sp['shot'], 'l'))
    out.append(text(9.35, 1.6 + sh + 0.08, 3.4, 0.4, [P(sp['shot_caption'], 14, GREEN, after=0)]))
    out.append(rect(9.35, 3.1, 3.25, 2.3))
    out.append(rect(9.35, 3.1, 0.07, 2.3, fill=AMBER))
    out.append(text(9.6, 3.1, 2.9, 2.3, [P(t, 18, TEXT, True, after=12) for t in sp['callout']], anchor='m'))
    return out


def L_bars(sp, badge):
    out = chrome(sp, badge)
    top = 1.95
    total = sum(v for _, v in sp['bars'])
    out.append(rect(0.7, top, 7.9, 4.75))
    out.append(text(1.0, top + 0.2, 7.4, 0.4, [P(sp['chart_title'], 15, TEXT, True, after=0)]))
    out.append(text(1.0, top + 0.55, 7.4, 0.35, [P('Share of the 256 values the uncopied byte $CD can hold', 12, MUTED, after=0)]))
    lab_x, bar_x, bar_max = 1.0, 2.45, 3.9
    out.append(text(lab_x, top + 1.0, 1.4, 0.3, [P('Loop length', 11, MUTED, after=0)]))
    vmax = max(v for _, v in sp['bars'])
    y = top + 1.4
    rows_h = 0.62
    out.append(rect(bar_x - 0.01, y - 0.12, 0.012, rows_h * len(sp['bars']), fill=RULE))
    for label, v in sp['bars']:
        is_hl = label == sp['hl_bar']
        bw = max(0.04, bar_max * v / vmax)
        pct = 100.0 * v / total
        tip = '%s-number loop: %d of %d cold starts (%.1f%%)' % (label, v, total, pct)
        out.append(text(lab_x, y - 0.06, 1.35, 0.36, [P(label, 15, AMBER if is_hl else TEXT, is_hl, MONO, 0)], 'r', 'm'))
        out.append(rect(bar_x, y, bw, 0.24, fill=AMBER if is_hl else BARGRAY, tip=tip))
        out.append(text(bar_x + bw + 0.12, y - 0.06, 2.2, 0.36,
                        [P('%.1f%%  (%d)' % (pct, v), 13, TEXT if is_hl else MUTED, is_hl, after=0)], anchor='m'))
        y += rows_h
    out += side_panels(8.85, top, 3.75, sp['side'], 4.75)
    return out


def L_code(sp, badge):
    out = chrome(sp, badge)
    top = 1.35
    out.append(rect(0.7, top, 7.9, 5.4))
    out.append(text(0.95, top + 0.2, 7.5, 5.1, [P(t, 10.5, TEXT, font=MONO, after=0) for t in sp['code']]))
    out += side_panels(8.85, top, 3.75, sp['side'], 5.4)
    return out


def L_quote(sp, badge):
    out = chrome(sp, badge)
    out.append(rect(0.7, 1.75, 11.9, 2.9))
    out.append(rect(0.7, 1.75, 0.07, 2.9, fill=AMBER))
    out.append(text(1.2, 1.85, 11.0, 2.2, [P('“' + sp['quote'] + '”', 26, TEXT, after=0)], anchor='m'))
    out.append(text(1.2, 4.0, 11.0, 0.5, [P(sp['attribution'], 15, MUTED, after=0)]))
    out.append(text(0.9, 5.0, 11.5, 1.5, [P('▸  ' + t, 19, TEXT, after=10) for t in sp['lines']]))
    return out


def L_memmap(sp, badge):
    out = chrome(sp, badge)
    y = 1.95
    heights = {'copy': 0.78, 'keep': 0.5, 'patch': 0.7}
    for addr, label, kind in sp['regions']:
        h = heights[kind]
        outline = AMBER if kind == 'patch' else None
        out.append(rect(0.7, y, 7.9, h - 0.06, fill=PANEL if kind != 'keep' else '1B1C1F', outline=outline))
        out.append(text(0.95, y, 2.3, h - 0.06, [P(addr, 14, AMBER if kind == 'patch' else GREEN, True, MONO, 0)], anchor='m'))
        out.append(text(3.3, y, 5.2, h - 0.06, [P(label, 14, TEXT if kind != 'keep' else MUTED, kind == 'patch', after=0)], anchor='m'))
        y += h
    out.append(text(0.7, y + 0.05, 7.9, 0.35, [P('Language card RAM, top to bottom. Not to scale.', 11, MUTED, after=0)]))
    out += side_panels(8.85, 1.95, 3.75, sp['side'], 4.6)
    return out


def L_closing(sp, badge):
    out = [rect(11.95, 0.42, 0.62, 0.62), text(11.95, 0.42, 0.62, 0.62, [P(badge, 15, GREEN, True, MONO, 0)], 'c', 'm')]
    out.append(text(0.9, 0.75, 10.6, 1.7, [P(t, 19, TEXT, after=8) for t in sp['takeaways']]))
    out.append(text(0.9, 2.55, 11.5, 0.8, [P(sp['title'], 44, TEXT, True, after=0)]))
    out.append(rect(0.92, 3.5, 1.6, 0.03, fill=GREEN))
    out.append(text(0.9, 3.8, 8.6, 0.5, [P('[ URL: Jeff to supply ]', 20, AMBER, font=MONO, after=0)]))
    out.append(rect(10.3, 3.6, 2.0, 2.0, outline=AMBER, dash=True))
    out.append(text(10.3, 3.6, 2.0, 2.0, [P('QR', 16, AMBER, True, after=2), P('Jeff to supply', 12, AMBER, after=0)], 'c', 'm'))
    out.append(text(0.9, 5.9, 9.0, 0.4, [P(sp['byline'], 15, MUTED, after=0)]))
    return out


LAYOUTS = dict(title=L_title, hero=L_hero, cols=L_cols, bullets=L_bullets, table=L_table, split=L_split,
               image=L_image, manual=L_manual, bars=L_bars, code=L_code, quote=L_quote, memmap=L_memmap,
               closing=L_closing)


def number_slides():
    n_main = n_backup = 0
    for sp in SLIDES:
        if sp['section'] == 'main':
            sp['badge'] = str(n_main)
            n_main += 1
        else:
            n_backup += 1
            sp['badge'] = 'B%d' % n_backup
        sp['prims'] = LAYOUTS[sp['kind']](sp, sp['badge'])


# ================================================================ renderers
def rgb(h):
    return RGBColor.from_string(h)


def render_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for sp in SLIDES:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BG)
        for pr in sp['prims']:
            kind = pr[0]
            if kind == 'rect':
                _, x, y, w, h, fill, outline, dash, _tip = pr
                shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(fill)
                if outline:
                    shp.line.color.rgb = rgb(outline)
                    shp.line.width = Pt(1.5)
                    if dash:
                        shp.line.dash_style = MSO_LINE.DASH
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
            elif kind == 'text':
                _, x, y, w, h, paras, align, anchor = pr
                tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE}[anchor]
                for i, pa in enumerate(paras):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
                    p.space_after = Pt(pa['after'])
                    r = p.add_run()
                    r.text = pa['t']
                    f = r.font
                    f.name, f.size, f.bold = pa['font'], Pt(pa['size']), pa['bold']
                    f.color.rgb = rgb(pa['color'])
            elif kind == 'img':
                _, x, y, w, h, asset, align = pr
                iw, ih = fit(asset, w, h)
                ix = x + (w - iw) / 2 if align == 'c' else x
                pic = s.shapes.add_picture(os.path.join(ASSETS, asset), Inches(ix), Inches(y), Inches(iw), Inches(ih))
                pic._element.nvPicPr.cNvPr.set('descr', asset.replace('_', ' ').replace('.png', ''))
        s.notes_slide.notes_text_frame.text = sp['notes']
    cp = prs.core_properties
    cp.title = "Applesoft's Loaded Dice (expanded edition)"
    cp.author = 'Draft for Jeff Robison'
    cp.subject = 'The Truth About Randomness on the Apple II, Then and Now'
    prs.save(path)


def data_uri(asset):
    with open(os.path.join(ASSETS, asset), 'rb') as fh:
        return 'data:image/png;base64,' + base64.b64encode(fh.read()).decode()


def render_html(path):
    px = 96.0
    sections = []
    for idx, sp in enumerate(SLIDES):
        parts = []
        for pr in sp['prims']:
            kind = pr[0]
            if kind == 'rect':
                _, x, y, w, h, fill, outline, dash, tip = pr
                border = 'border:%.1fpx %s #%s;' % (2, 'dashed' if dash else 'solid', outline) if outline else ''
                cls = ' class="bar" tabindex="0"' if tip else ''
                title = ' title="%s" aria-label="%s"' % (html.escape(tip), html.escape(tip)) if tip else ''
                parts.append('<div%s%s style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;background:#%s;%s"></div>'
                             % (cls, title, x * px, y * px, w * px, h * px, fill, border))
            elif kind == 'text':
                _, x, y, w, h, paras, align, anchor = pr
                ps = []
                for pa in paras:
                    fam = "'Courier New',Courier,monospace" if pa['font'] == MONO else 'Arial,Helvetica,sans-serif'
                    ps.append('<p style="font-size:%.1fpx;color:#%s;font-weight:%s;font-family:%s;margin:0 0 %.1fpx 0;%s">%s</p>'
                              % (pa['size'] * 96 / 72, pa['color'], 700 if pa['bold'] else 400, fam,
                                 pa['after'] * 96 / 72, 'white-space:pre-wrap;' if pa['font'] == MONO else '',
                                 html.escape(pa['t'])))
                parts.append('<div class="t" style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;text-align:%s;justify-content:%s">%s</div>'
                             % (x * px, y * px, w * px, h * px, {'l': 'left', 'c': 'center', 'r': 'right'}[align],
                                {'t': 'flex-start', 'm': 'center'}[anchor], ''.join(ps)))
            elif kind == 'img':
                _, x, y, w, h, asset, align = pr
                iw, ih = fit(asset, w, h)
                ix = x + (w - iw) / 2 if align == 'c' else x
                parts.append('<img alt="%s" src="%s" style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx">'
                             % (html.escape(asset), data_uri(asset), ix * px, y * px, iw * px, ih * px))
        sections.append('<section class="slide" data-i="%d" data-badge="%s">%s<aside class="notes">%s</aside></section>'
                        % (idx, sp['badge'], ''.join(parts), html.escape(sp['notes']).replace('\n', '<br>')))
    doc = HTML_TEMPLATE.replace('{{SLIDES}}', '\n'.join(sections)).replace('{{COUNT}}', str(len(SLIDES)))
    with open(path, 'w') as fh:
        fh.write(doc)


HTML_TEMPLATE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>Applesoft's Loaded Dice (expanded edition)</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
html,body{margin:0;height:100%;background:#0e0f10;color:#E9E7E3;font-family:Arial,Helvetica,sans-serif;overflow:hidden}
#stage{position:absolute;left:50%;top:50%;width:1280px;height:720px;transform-origin:0 0}
.slide{position:absolute;inset:0;width:1280px;height:720px;background:#17181A;display:none;overflow:hidden}
.slide.on{display:block}
.slide>div,.slide>img{position:absolute;box-sizing:border-box}
.slide .t{display:flex;flex-direction:column;line-height:1.2}
.slide .bar{border-radius:0 4px 4px 0;cursor:default}
.slide .bar:hover,.slide .bar:focus{outline:2px solid #E9E7E3;outline-offset:2px}
.notes{display:none}
#notes{position:fixed;left:0;right:0;bottom:0;max-height:34vh;overflow:auto;background:#202225;color:#E9E7E3;
  padding:14px 22px;font-size:15px;line-height:1.45;border-top:1px solid #3A3E42;display:none}
#hud{position:fixed;right:14px;bottom:10px;font:12px 'Courier New',monospace;color:#9DA39B}
body.shownotes #notes{display:block}
body.print{overflow:auto}
body.print #stage{position:static;transform:none!important;margin:0 auto}
body.print .slide{position:relative;display:block;margin:0 auto 24px}
body.print #hud{display:none}
@media print{body{background:#fff}#hud,#notes{display:none!important}
  #stage{position:static;transform:none!important}.slide{display:block;position:relative;page-break-after:always}}
</style></head>
<body>
<div id="stage">
{{SLIDES}}
</div>
<div id="notes"></div>
<div id="hud"></div>
<script>
(function(){
  var slides=[].slice.call(document.querySelectorAll('.slide')), i=0, n={{COUNT}};
  var stage=document.getElementById('stage'), notes=document.getElementById('notes'), hud=document.getElementById('hud');
  function fit(){ if(document.body.classList.contains('print'))return;
    var s=Math.min(window.innerWidth/1280, window.innerHeight/720);
    stage.style.transform='scale('+s+') translate(-50%,-50%)'; stage.style.transformOrigin='0 0';
    stage.style.left='50%'; stage.style.top='50%';
    stage.style.transform='translate(-50%,-50%) scale('+s+')'; stage.style.transformOrigin='center center'; }
  function show(k){ i=Math.max(0,Math.min(n-1,k)); slides.forEach(function(el,j){el.classList.toggle('on',j===i);});
    notes.innerHTML=slides[i].querySelector('.notes').innerHTML;
    hud.textContent=slides[i].dataset.badge+'  \\u00b7  '+(i+1)+'/'+n+'   \\u2190 \\u2192  N notes  P all';
    if(location.hash!=='#'+(i+1)) history.replaceState(null,'','#'+(i+1)); }
  document.addEventListener('keydown',function(e){
    if(['ArrowRight','PageDown',' '].indexOf(e.key)>=0){show(i+1);e.preventDefault();}
    else if(['ArrowLeft','PageUp'].indexOf(e.key)>=0){show(i-1);e.preventDefault();}
    else if(e.key==='Home')show(0); else if(e.key==='End')show(n-1);
    else if(e.key==='n'||e.key==='N')document.body.classList.toggle('shownotes');
    else if(e.key==='p'||e.key==='P'){document.body.classList.toggle('print');
      slides.forEach(function(el){el.classList.add('on');}); if(!document.body.classList.contains('print'))show(i); fit();}
  });
  document.addEventListener('click',function(e){ if(e.target.closest('#notes'))return;
    if(document.body.classList.contains('print'))return; show(e.clientX>window.innerWidth/3?i+1:i-1); });
  window.addEventListener('resize',fit);
  if(/[?&]print/.test(location.search)){document.body.classList.add('print');slides.forEach(function(el){el.classList.add('on');});}
  fit(); show((parseInt(location.hash.slice(1),10)||1)-1);
})();
</script>
</body></html>
"""


def screen_lines(sp):
    """Everything the audience reads on a slide, as plain lines (for Marp and the outline)."""
    k = sp['kind']
    out = []
    if sp.get('subtitle'):
        out += [sp['subtitle'], sp['byline']]
    if sp.get('lede'):
        out.append(sp['lede'])
    if k == 'hero':
        out += ['`%s`' % t for t in sp['mono'] if t] + sp['sub']
    if k == 'cols':
        for col in sp['cols']:
            out.append('**%s:** %s' % (col[0], ' / '.join(t.strip('`') for t in col[1])))
    if k == 'bullets':
        out += sp['bullets']
    if k in ('table', 'split'):
        out.append(' | '.join(h or ' ' for h in sp['header']))
        out += [' | '.join(r) for r in sp['rows']]
    if k in ('split', 'code'):
        out += ['`%s`' % t for t in sp['code'] if t]
    if k in ('image', 'bars', 'memmap', 'code'):
        for item in sp.get('side', []):
            out.append('**%s:** %s' % (item[0], ' / '.join(item[1])))
    if k == 'image':
        out.append('_%s_' % sp['caption'])
    if k == 'manual':
        out += ['%s (manual scan)' % r[0] for r in sp['rows']] + [sp['shot_caption']] + sp['callout']
    if k == 'bars':
        total = sum(v for _, v in sp['bars'])
        out.append(sp['chart_title'])
        out += ['%s-number loop: %d of %d (%.1f%%)' % (l, v, total, 100.0 * v / total) for l, v in sp['bars']]
    if k == 'memmap':
        out += ['%s  %s' % (a, l) for a, l, _ in sp['regions']]
    if k == 'quote':
        out += ['“%s”' % sp['quote'], sp['attribution']] + sp['lines']
    if k == 'closing':
        out += sp['takeaways'] + ['[ URL: Jeff to supply ]', '[ QR: Jeff to supply ]', sp['byline']]
    c = sp.get('callout')
    if c and k != 'manual':
        out.append('> ' + c)
    return out


def assets_of(sp):
    a = []
    if sp.get('image'):
        a.append(sp['image'])
    if sp['kind'] == 'manual':
        a += [r[1] for r in sp['rows']] + [sp['shot']]
    return a


def render_marp(path):
    head = """---
marp: true
size: 16:9
paginate: false
title: "Applesoft's Loaded Dice (expanded edition)"
style: |
  section { background:#17181A; color:#E9E7E3; font-family:Arial,Helvetica,sans-serif; font-size:26px; padding:48px 64px; }
  h1 { color:#E9E7E3; font-size:46px; margin:0 0 12px 0; }
  h2 { color:#5FCB7E; font-size:26px; font-weight:400; }
  strong { color:#5FCB7E; }
  code { background:#202225; color:#5FCB7E; font-family:'Courier New',monospace; }
  blockquote { border-left:6px solid #E0A73E; background:#202225; color:#E9E7E3; padding:12px 20px; }
  table { font-size:20px; } th { color:#9DA39B; background:#17181A; } td { background:#202225; }
  section::after { content: attr(data-badge); position:absolute; top:28px; right:40px; color:#5FCB7E; font-family:'Courier New',monospace; }
  img { background:transparent; }
  footer { color:#9DA39B; font-size:14px; }
---
"""
    chunks = []
    for sp in SLIDES:
        md = ['<!-- _footer: "%s" -->' % sp.get('source', '').replace('"', "'")] if sp.get('source') else []
        md.append('<!-- _class: %s -->' % sp['section'])
        md.append('# %s%s' % ('' if sp['section'] == 'main' else '', sp['title']))
        k = sp['kind']
        if k in ('table',):
            if sp.get('lede'):
                md.append('## ' + sp['lede'])
            md.append('| ' + ' | '.join(h or ' ' for h in sp['header']) + ' |')
            md.append('|' + '---|' * len(sp['header']))
            md += ['| ' + ' | '.join(c or ' ' for c in r) + ' |' for r in sp['rows']]
            if sp.get('callout'):
                md.append('\n> ' + sp['callout'])
        elif k in ('split', 'code'):
            if sp.get('lede'):
                md.append('## ' + sp['lede'])
            md.append('```\n' + '\n'.join(sp['code']) + '\n```')
            if k == 'split':
                md.append('| ' + ' | '.join(sp['header']) + ' |\n|---|---|')
                md += ['| ' + ' | '.join(r) + ' |' for r in sp['rows']]
                md.append('\n> ' + sp['callout'])
            else:
                md += ['**%s:** %s' % (i[0], ' · '.join(i[1])) for i in sp['side']]
        else:
            for a in assets_of(sp):
                md.append('![w:%d](assets/%s)' % (560 if 'manual' in a else 700, a))
            for ln in screen_lines(sp):
                if ln.startswith('> ') or ln.startswith('**') or ln.startswith('_'):
                    md.append(ln)
                elif ln.startswith('`'):
                    md.append(ln)
                else:
                    md.append('- ' + ln if k in ('bullets', 'bars', 'memmap', 'closing', 'manual') else ln)
                md.append('')
        md.append('\n<!--\n%s\n-->' % sp['notes'])
        chunks.append('<!-- badge: %s -->\n' % sp['badge'] + '\n'.join(md))
    with open(path, 'w') as fh:
        fh.write(head + '\n\n---\n\n'.join(chunks) + '\n')


def render_outline(path):
    lines = ["# Keynote outline: *Applesoft's Loaded Dice* (expanded edition)", '',
             'Per-slide outline for building or adjusting the deck in Keynote. '
             '`Applesoft_Loaded_Dice_EXPANDED.pptx` opens directly in Keynote; this file is the same content as '
             'text, with assets and sources per slide. Generated from `_build/build_new_deck.py`.', '',
             'Palette: background `#17181A`, panels `#202225`, accent green `#5FCB7E`, text `#E9E7E3`, '
             'muted `#9DA39B`, highlight amber `#E0A73E`. Fonts: Arial, Courier New.', '',
             'Main talk: slides 0–%d. Backup: B1–B%d.' % (
                 sum(1 for s in SLIDES if s['section'] == 'main') - 1,
                 sum(1 for s in SLIDES if s['section'] == 'backup')), '']
    for sp in SLIDES:
        lines.append('---\n')
        lines.append('## %s. %s' % (sp['badge'], sp['title']))
        lines.append('')
        lines.append('**On screen**')
        lines.append('')
        for ln in screen_lines(sp):
            lines.append('- ' + ln.lstrip('> '))
        a = assets_of(sp)
        if a:
            lines.append('')
            lines.append('**Assets:** ' + ', '.join('`assets/%s`' % x for x in a))
        if sp.get('source'):
            lines.append('')
            lines.append('**Source line:** ' + sp['source'])
        lines.append('')
        lines.append('**Speaker notes**')
        lines.append('')
        lines += ['> ' + ln if ln else '>' for ln in sp['notes'].split('\n')]
        lines.append('')
    with open(path, 'w') as fh:
        fh.write('\n'.join(lines))


def main():
    number_slides()
    pptx_path = os.path.join(OUTDIR, BASENAME + '.pptx')
    render_pptx(pptx_path)
    render_html(os.path.join(OUTDIR, BASENAME + '.html'))
    render_marp(os.path.join(OUTDIR, BASENAME + '_marp.md'))
    render_outline(os.path.join(OUTDIR, 'KEYNOTE_OUTLINE.md'))
    prs = Presentation(pptx_path)
    print('slides:', len(prs.slides), ' badges:', [sp['badge'] for sp in SLIDES])
    for name in os.listdir(OUTDIR):
        p = os.path.join(OUTDIR, name)
        if os.path.isfile(p):
            print('%9d  %s' % (os.path.getsize(p), name))


if __name__ == '__main__':
    main()
